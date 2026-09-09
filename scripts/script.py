#!/usr/bin/env python3
"""Recover requested Badge/VIS documents from barcodes in lot PDFs.

Each worker renders exactly one PDF page at 300 DPI, copies only its upper
half into a new image, releases the full-page render, and gives that upper-half
image to CodaraScan's Panorama engine.  Workers never write shared outputs;
the parent process performs matching, reporting, CSV writing, and workbook
highlighting.
"""

from __future__ import annotations

import argparse
import atexit
import csv
import logging
import multiprocessing as mp
import os
import re
import sys
import tempfile
import time
import traceback
import unicodedata
from collections import defaultdict
from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, TextIO

DPI = 300
PDF_POINTS_PER_INCH = 72
PANORAMA_MODE = "panorama"
PANORAMA_ENGINE = "panorama-extractor"
BARCODE_SYMBOLS = "linear"
BARCODE_FORMATS = ("code-39", "code-128")
CSV_NAME = "results.csv"
REPORT_NAME = "report.txt"
LOG_NAME = "processing.log"
GREEN_FILL = "C6EFCE"

# Badges are numeric identifiers containing between one and five digits.
# Keep them as strings so leading zeroes remain significant when present.
BADGE_RE = re.compile(r"^\d{1,5}$")
# ISO 3779 VIN characters. I, O and Q are excluded because they are easily
# confused with 1 and 0. A check-digit test is intentionally not imposed:
# it is not mandatory in every market represented by this dataset.
VIN_RE = re.compile(r"^[A-HJ-NPR-Z0-9]{17}$")
VIS_RE = re.compile(r"^[A-HJ-NPR-Z0-9]{8}$")


class ConfigurationError(RuntimeError):
    """A fatal input or runtime configuration problem."""


@dataclass(frozen=True, order=True)
class RowRef:
    sheet: str
    row: int


@dataclass(frozen=True)
class TargetRow:
    ref: RowRef
    badge: str | None
    vin: str | None
    vis: str | None


@dataclass
class TargetIndex:
    workbook_path: Path
    rows: list[TargetRow] = field(default_factory=list)
    badge_rows: dict[str, set[RowRef]] = field(default_factory=dict)
    vis_rows: dict[str, set[RowRef]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class LotPlan:
    lot: str
    pdf_path: Path
    pages: int


@dataclass(frozen=True)
class PageTask:
    lot: str
    pdf_path: str
    page_index: int
    pages_in_lot: int


@dataclass(frozen=True)
class PageResult:
    lot: str
    page_number: int
    pages_in_lot: int
    badges: tuple[str, ...] = ()
    vins: tuple[str, ...] = ()
    engine: str = ""
    mode: str = ""
    error: str | None = None
    traceback_text: str | None = None


@dataclass(frozen=True)
class CsvRecord:
    lot: str
    page_number: int
    badge: str = ""
    vin: str = ""
    vis: str = ""
    found: str = ""


@dataclass
class RunStats:
    discovered_lots: int = 0
    processed_lots: int = 0
    processed_pdfs: int = 0
    planned_pages: int = 0
    completed_pages: int = 0
    successful_pages: int = 0
    failed_pages: int = 0
    badge_matched_rows: set[RowRef] = field(default_factory=set)
    vis_matched_rows: set[RowRef] = field(default_factory=set)
    errors: list[str] = field(default_factory=list)


_WORKER_SCANNER = None
_WORKER_PDF = None
_WORKER_PDF_PATH: str | None = None


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Render each lot PDF page at 300 DPI, scan its upper half with "
            "CodaraScan Panorama, and highlight matching Excel rows."
        )
    )
    parser.add_argument(
        "-d", "--directory", required=True, type=Path, help="Batch root directory"
    )
    parser.add_argument(
        "-n",
        "--workers",
        required=True,
        type=int,
        help="Number of page worker processes",
    )
    parser.add_argument(
        "--excel",
        type=Path,
        help="Source workbook path or filename (needed only if discovery is ambiguous)",
    )
    args = parser.parse_args(argv)
    if args.workers <= 0:
        parser.error("--workers/-n must be greater than zero")
    return args


def configure_logging(root: Path) -> logging.Logger:
    logger = logging.getLogger("digitization_recovery")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    logger.propagate = False
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    handler = logging.FileHandler(root / LOG_NAME, mode="w", encoding="utf-8")
    handler.setFormatter(formatter)
    logger.addHandler(handler)
    return logger


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return "".join(character for character in text.casefold() if character.isalnum())


BADGE_HEADERS = {"badge", "bdg", "nobadge", "numerobadge", "numbadge"}
VIN_HEADERS = {"vin", "novin", "numerovin", "numvin"}
VIS_HEADERS = {"vis", "novis", "numerovis", "numvis"}


def normalize_badge(value: object, number_format: str = "General") -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        text = str(value)
        if len(text) < 5 and "00000" in number_format:
            text = f"{value:05d}"
    elif isinstance(value, float) and value.is_integer():
        integer = int(value)
        text = str(integer)
        if len(text) < 5 and "00000" in number_format:
            text = f"{integer:05d}"
    else:
        text = str(value).strip()
    return text if BADGE_RE.fullmatch(text) else None


def normalize_vin(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper()
    return text if VIN_RE.fullmatch(text) else None


def normalize_vis(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper()
    return text if VIS_RE.fullmatch(text) else None


def classify_decoded_value(value: object) -> tuple[str, str] | None:
    """Return (kind, normalized value) for an exact Badge or VIN payload."""

    if value is None:
        return None
    text = str(value).strip().upper()
    if BADGE_RE.fullmatch(text):
        return "badge", text
    if VIN_RE.fullmatch(text):
        return "vin", text
    return None


def _header_columns(
    worksheet: Any,
) -> tuple[int, int, int | None, int | None] | None:
    max_search_row = min(getattr(worksheet, "max_row"), 50)
    max_search_column = min(getattr(worksheet, "max_column"), 100)
    for row_number in range(1, max_search_row + 1):
        badge_column = None
        vin_column = None
        vis_column = None
        for column in range(1, max_search_column + 1):
            header = normalize_header(worksheet.cell(row_number, column).value)
            if header in BADGE_HEADERS and badge_column is None:
                badge_column = column
            elif header in VIN_HEADERS and vin_column is None:
                vin_column = column
            elif header in VIS_HEADERS and vis_column is None:
                vis_column = column
        if badge_column is not None and (
            vin_column is not None or vis_column is not None
        ):
            return row_number, badge_column, vin_column, vis_column
    return None


def discover_excel(root: Path, explicit: Path | None) -> Path:
    if explicit is not None:
        candidate = explicit if explicit.is_absolute() else root / explicit
        candidate = candidate.resolve()
        if candidate.parent != root:
            raise ConfigurationError(
                "Source workbook must be directly inside the batch root"
            )
        if not candidate.is_file():
            raise ConfigurationError(f"Excel workbook does not exist: {candidate}")
        if candidate.suffix.casefold() not in {".xlsx", ".xlsm"}:
            raise ConfigurationError("Source workbook must be .xlsx or .xlsm")
        return candidate

    candidates = sorted(
        path
        for path in root.iterdir()
        if path.is_file()
        and path.suffix.casefold() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
        and not path.stem.casefold().endswith("_found")
    )
    if not candidates:
        raise ConfigurationError(f"No .xlsx or .xlsm source workbook found in {root}")
    if len(candidates) > 1:
        names = ", ".join(path.name for path in candidates)
        raise ConfigurationError(
            f"Multiple source workbooks found ({names}); select one with --excel"
        )
    return candidates[0]


def load_target_index(workbook_path: Path) -> TargetIndex:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConfigurationError(
            "openpyxl is required; install dependencies with "
            "'python -m pip install -r requirements.txt'"
        ) from exc

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_links=True,
            keep_vba=workbook_path.suffix.casefold() == ".xlsm",
        )
    except Exception as exc:
        raise ConfigurationError(
            f"Cannot open Excel workbook {workbook_path}: {exc}"
        ) from exc

    index = TargetIndex(workbook_path=workbook_path)
    badge_map: defaultdict[str, set[RowRef]] = defaultdict(set)
    vis_map: defaultdict[str, set[RowRef]] = defaultdict(set)
    recognized_sheets = 0
    try:
        for worksheet in workbook.worksheets:
            columns = _header_columns(worksheet)
            if columns is None:
                continue
            recognized_sheets += 1
            header_row, badge_column, vin_column, vis_column = columns
            for row_number in range(header_row + 1, worksheet.max_row + 1):
                badge_cell = worksheet.cell(row_number, badge_column)
                vin_cell = (
                    worksheet.cell(row_number, vin_column) if vin_column else None
                )
                vis_cell = (
                    worksheet.cell(row_number, vis_column) if vis_column else None
                )
                raw_values = (
                    badge_cell.value,
                    vin_cell.value if vin_cell else None,
                    vis_cell.value if vis_cell else None,
                )
                if all(
                    value is None or str(value).strip() == "" for value in raw_values
                ):
                    continue

                ref = RowRef(worksheet.title, row_number)
                badge = normalize_badge(badge_cell.value, badge_cell.number_format)
                vin = normalize_vin(vin_cell.value) if vin_cell else None
                explicit_vis = normalize_vis(vis_cell.value) if vis_cell else None
                derived_vis = vin[-8:] if vin else None
                vis = explicit_vis or derived_vis

                problems: list[str] = []
                if badge_cell.value not in (None, "") and badge is None:
                    problems.append(f"invalid Badge {badge_cell.value!r}")
                if vin_cell and vin_cell.value not in (None, "") and vin is None:
                    problems.append(f"invalid VIN {vin_cell.value!r}")
                if (
                    vis_cell
                    and vis_cell.value not in (None, "")
                    and explicit_vis is None
                ):
                    problems.append(f"invalid VIS {vis_cell.value!r}")
                if explicit_vis and derived_vis and explicit_vis != derived_vis:
                    problems.append(
                        f"VIS {explicit_vis!r} conflicts with VIN-derived VIS "
                        f"{derived_vis!r}"
                    )
                    # An inconsistent row may still be recovered by Badge, but
                    # neither conflicting vehicle identifier is safe to index.
                    vis = None
                if problems:
                    index.warnings.append(
                        f"Excel {worksheet.title}!{row_number}: " + "; ".join(problems)
                    )

                target = TargetRow(ref=ref, badge=badge, vin=vin, vis=vis)
                index.rows.append(target)
                if badge:
                    badge_map[badge].add(ref)
                if vis:
                    vis_map[vis].add(ref)
                if not badge and not vis:
                    index.warnings.append(
                        f"Excel {worksheet.title}!{row_number}: "
                        "no usable Badge or VIS/VIN"
                    )
    finally:
        workbook.close()

    if recognized_sheets == 0:
        raise ConfigurationError(
            "No worksheet contains a Badge/BDG column plus a VIN or VIS column"
        )
    if not index.rows:
        raise ConfigurationError("The source workbook contains no requested data rows")
    if not badge_map and not vis_map:
        raise ConfigurationError(
            "The source workbook contains no valid Badge or VIS/VIN values"
        )
    index.badge_rows = dict(badge_map)
    index.vis_rows = dict(vis_map)
    return index


def lot_sort_key(name: str) -> tuple[int, int | str, str]:
    stripped = name.strip()
    if stripped.isdecimal():
        return 0, int(stripped), stripped
    return 1, stripped.casefold(), stripped


def discover_lots(
    root: Path, logger: logging.Logger
) -> tuple[list[LotPlan], list[str], int]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise ConfigurationError(
            "pypdfium2 is required (it is installed with CodaraScan)"
        ) from exc

    directories = sorted(
        (path for path in root.iterdir() if path.is_dir()),
        key=lambda path: lot_sort_key(path.name),
    )
    if not directories:
        raise ConfigurationError(f"No immediate Lot subdirectories found in {root}")

    plans: list[LotPlan] = []
    errors: list[str] = []
    for directory in directories:
        pdfs = sorted(
            path
            for path in directory.iterdir()
            if path.is_file() and path.suffix.casefold() == ".pdf"
        )
        if len(pdfs) != 1:
            detail = (
                "no PDF"
                if not pdfs
                else f"multiple PDFs ({', '.join(p.name for p in pdfs)})"
            )
            message = f"Lot {directory.name}: {detail}; lot skipped"
            errors.append(message)
            logger.error(message)
            continue
        document = None
        try:
            document = pdfium.PdfDocument(str(pdfs[0]))
            page_count = len(document)
            if page_count <= 0:
                raise ValueError("PDF contains no pages")
            plans.append(LotPlan(directory.name, pdfs[0], page_count))
        except Exception as exc:
            message = (
                f"Lot {directory.name}: cannot open {pdfs[0].name}: "
                f"{type(exc).__name__}: {exc}"
            )
            errors.append(message)
            logger.exception(message)
        finally:
            if document is not None:
                try:
                    document.close()
                except Exception:
                    pass
    if not plans:
        raise ConfigurationError("No processable Lot PDFs were found")
    return plans, errors, len(directories)


def create_panorama_scanner() -> Any:
    """Create the one supported scanner configuration for this dataset."""

    from codarascan import Scanner

    return Scanner(
        mode=PANORAMA_MODE,
        symbols=BARCODE_SYMBOLS,
        formats=BARCODE_FORMATS,
        decode=True,
    )


def validate_runtime() -> None:
    try:
        import codarascan
    except ImportError as exc:
        raise ConfigurationError(
            "CodaraScan is required; install dependencies with "
            "'python -m pip install -r requirements.txt'"
        ) from exc

    version = getattr(codarascan, "__version__", "unknown")
    scanner = create_panorama_scanner()
    try:
        scanner.warm()
    except Exception as exc:
        raise ConfigurationError(
            f"CodaraScan Panorama failed to initialize: {exc}"
        ) from exc
    actual_mode = getattr(scanner.mode, "value", scanner.mode)
    if actual_mode != PANORAMA_MODE:
        raise ConfigurationError(
            f"CodaraScan {version} did not retain panorama mode (got {actual_mode!r})"
        )


def initialize_worker() -> None:
    global _WORKER_SCANNER
    # Keep native math/OpenCV pools from multiplying the requested process
    # count. These defaults are set before importing CodaraScan in each spawned
    # worker; CodaraScan still performs its own Panorama stages normally.
    os.environ["OMP_NUM_THREADS"] = "1"
    os.environ["OPENBLAS_NUM_THREADS"] = "1"
    os.environ["MKL_NUM_THREADS"] = "1"
    try:
        import cv2

        cv2.setNumThreads(1)
    except Exception:
        pass
    _WORKER_SCANNER = create_panorama_scanner()
    _WORKER_SCANNER.warm()
    atexit.register(_close_worker_document)


def _close_worker_document() -> None:
    global _WORKER_PDF, _WORKER_PDF_PATH
    if _WORKER_PDF is not None:
        try:
            _WORKER_PDF.close()
        except Exception:
            pass
    _WORKER_PDF = None
    _WORKER_PDF_PATH = None


def _worker_document(pdf_path: str) -> Any:
    global _WORKER_PDF, _WORKER_PDF_PATH
    if _WORKER_PDF is not None and _WORKER_PDF_PATH == pdf_path:
        return _WORKER_PDF
    _close_worker_document()
    import pypdfium2 as pdfium

    _WORKER_PDF = pdfium.PdfDocument(pdf_path)
    _WORKER_PDF_PATH = pdf_path
    return _WORKER_PDF


def render_upper_half(document: Any, page_index: int) -> Any:
    """Render a full page at 300 DPI, then return an independent top-half PIL image."""

    page = None
    bitmap = None
    full_image = None
    try:
        page = document[page_index]
        # PDF coordinates use 72 points/inch, hence 300/72 is exactly 300 DPI.
        bitmap = page.render(scale=DPI / PDF_POINTS_PER_INCH, rotation=0)
        full_image = bitmap.to_pil()
        width, height = full_image.size
        if width <= 0 or height < 2:
            raise ValueError(f"rendered page has invalid dimensions {width}x{height}")
        # Pillow.crop() creates an independent image. Closing the full image and
        # Pdfium bitmap below ensures the lower half cannot reach CodaraScan.
        return full_image.crop((0, 0, width, height // 2))
    finally:
        if full_image is not None:
            full_image.close()
        if bitmap is not None:
            try:
                bitmap.close()
            except Exception:
                pass
        if page is not None:
            try:
                page.close()
            except Exception:
                pass


def process_page(task: PageTask) -> PageResult:
    roi_image = None
    try:
        if _WORKER_SCANNER is None:
            raise RuntimeError("worker scanner was not initialized")
        document = _worker_document(task.pdf_path)
        roi_image = render_upper_half(document, task.page_index)
        # Deliberately do not pass CodaraScan's roi= parameter. Panorama receives
        # only the already-cropped upper-half pixels, never the full page.
        scan = _WORKER_SCANNER.scan_image(roi_image, diagnostics=False)
        mode = scan.metadata.mode
        engine = scan.metadata.engine
        if mode != PANORAMA_MODE or engine != PANORAMA_ENGINE:
            raise RuntimeError(
                "unexpected CodaraScan engine metadata: "
                f"mode={mode!r}, engine={engine!r}"
            )

        from codarascan import DecodedSymbolResult

        badges: set[str] = set()
        vins: set[str] = set()
        for symbol in scan.symbols:
            # This type check excludes localized, unresolved, and review-candidate
            # results even if future versions add payload-like attributes to them.
            if not isinstance(symbol, DecodedSymbolResult):
                continue
            classified = classify_decoded_value(symbol.text)
            if classified is None:
                continue
            kind, value = classified
            if kind == "badge":
                badges.add(value)
            else:
                vins.add(value)
        return PageResult(
            lot=task.lot,
            page_number=task.page_index + 1,
            pages_in_lot=task.pages_in_lot,
            badges=tuple(sorted(badges)),
            vins=tuple(sorted(vins)),
            engine=engine,
            mode=mode,
        )
    except Exception as exc:
        # Reopen the PDF on the next task in case a failed render left Pdfium's
        # document handle in an uncertain state.
        _close_worker_document()
        return PageResult(
            lot=task.lot,
            page_number=task.page_index + 1,
            pages_in_lot=task.pages_in_lot,
            error=f"{type(exc).__name__}: {exc}",
            traceback_text=traceback.format_exc(),
        )
    finally:
        if roi_image is not None:
            roi_image.close()


def iter_tasks(plans: Iterable[LotPlan]) -> Iterator[PageTask]:
    for plan in plans:
        for page_index in range(plan.pages):
            yield PageTask(plan.lot, str(plan.pdf_path), page_index, plan.pages)


def matched_rows_for_page(
    result: PageResult, targets: TargetIndex
) -> tuple[set[RowRef], set[RowRef]]:
    badge_matches: set[RowRef] = set()
    vis_matches: set[RowRef] = set()
    for badge in result.badges:
        badge_matches.update(targets.badge_rows.get(badge, ()))
    for vin in result.vins:
        vis_matches.update(targets.vis_rows.get(vin[-8:], ()))
    return badge_matches, vis_matches


def csv_records_for_page(result: PageResult, targets: TargetIndex) -> list[CsvRecord]:
    """Return the CSV representation of one processed page.

    Badge and VIN values decoded on that page stay on its single row. Missing
    values remain blank; nothing is inferred from the source workbook or from
    another page. A page with no identifiers still gets one empty row. Multiple
    distinct values are retained with semicolon separators as a defensive
    fallback, although the expected dataset contains at most one of each.
    """

    badge = ";".join(result.badges)
    vin = ";".join(result.vins)
    vis_values = tuple(decoded_vin[-8:] for decoded_vin in result.vins)
    vis = ";".join(vis_values)
    found = any(targets.badge_rows.get(value) for value in result.badges) or any(
        targets.vis_rows.get(value) for value in vis_values
    )
    return [
        CsvRecord(
            result.lot,
            result.page_number,
            badge=badge,
            vin=vin,
            vis=vis,
            found="FOUND" if found else "",
        )
    ]


def atomic_text_write(
    path: Path,
    writer: Callable[[TextIO], None],
    *,
    newline: str | None = None,
    encoding: str = "utf-8",
) -> None:
    file_descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent, prefix=f".{path.name}.", suffix=".tmp", text=True
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(
            file_descriptor, "w", encoding=encoding, newline=newline
        ) as handle:
            writer(handle)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, path)
    except Exception:
        try:
            temporary_path.unlink(missing_ok=True)
        finally:
            raise


def write_csv(path: Path, records: list[CsvRecord]) -> None:
    records.sort(
        key=lambda item: (
            lot_sort_key(item.lot),
            item.page_number,
            item.badge,
            item.vin,
        )
    )

    def write(handle: TextIO) -> None:
        csv_writer = csv.writer(handle)
        csv_writer.writerow(["Lot", "numero page", "bdg", "VIN", "VIS", "found"])
        for record in records:
            csv_writer.writerow(
                [
                    record.lot,
                    record.page_number,
                    record.badge,
                    record.vin,
                    record.vis,
                    record.found,
                ]
            )

    atomic_text_write(path, write, newline="", encoding="utf-8-sig")


def highlight_workbook(source: Path, output: Path, found_rows: set[RowRef]) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    keep_vba = source.suffix.casefold() == ".xlsm"
    workbook = load_workbook(
        source, data_only=False, keep_links=True, keep_vba=keep_vba
    )
    fill = PatternFill(fill_type="solid", fgColor=GREEN_FILL)
    for ref in sorted(found_rows):
        worksheet = workbook[ref.sheet]
        for column in range(1, worksheet.max_column + 1):
            worksheet.cell(ref.row, column).fill = fill

    file_descriptor, temporary_name = tempfile.mkstemp(
        dir=output.parent, prefix=f".{output.stem}.", suffix=output.suffix
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, output)
    except Exception:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
        raise


def format_duration(seconds: float) -> str:
    total = max(0, int(round(seconds)))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def write_report(
    path: Path,
    *,
    started_at: datetime,
    finished_at: datetime,
    elapsed: float,
    targets: TargetIndex,
    stats: RunStats,
    workbook_output: Path,
    csv_output: Path,
) -> None:
    found_rows = stats.badge_matched_rows | stats.vis_matched_rows
    both_rows = stats.badge_matched_rows & stats.vis_matched_rows
    total_rows = len(targets.rows)
    not_found = total_rows - len(found_rows)
    percentage = (100.0 * len(found_rows) / total_rows) if total_rows else 0.0
    all_issues = [*targets.warnings, *stats.errors]

    lines = [
        "Document digitization recovery report",
        "====================================",
        f"Start time: {started_at.isoformat(timespec='seconds')}",
        f"Finish time: {finished_at.isoformat(timespec='seconds')}",
        f"Total processing time: {format_duration(elapsed)}",
        "",
        f"Source workbook: {targets.workbook_path.name}",
        f"Updated workbook: {workbook_output.name}",
        f"CSV output: {csv_output.name}",
        "",
        f"Total requested rows: {total_rows:,}",
        f"Found: {len(found_rows):,}",
        f"Not found: {not_found:,}",
        f"Recovery rate: {percentage:.2f}%",
        f"Badge matches: {len(stats.badge_matched_rows):,}",
        f"VIS/VIN matches: {len(stats.vis_matched_rows):,}",
        f"Rows matched using both identifiers: {len(both_rows):,}",
        "",
        f"Lot directories discovered: {stats.discovered_lots:,}",
        f"Lots processed: {stats.processed_lots:,}",
        f"PDFs processed: {stats.processed_pdfs:,}",
        f"Total PDF pages planned: {stats.planned_pages:,}",
        f"Total pages processed: {stats.completed_pages:,}",
        f"Successfully processed pages: {stats.successful_pages:,}",
        f"Failed pages: {stats.failed_pages:,}",
        "",
        f"Errors and input warnings: {len(all_issues):,}",
    ]
    if all_issues:
        lines.extend(f"- {message}" for message in all_issues)
    else:
        lines.append("- None")
    lines.append("")

    def write(handle: TextIO) -> None:
        handle.write("\n".join(lines))

    atomic_text_write(path, write)


def run(args: argparse.Namespace) -> int:
    root = args.directory.expanduser().resolve()
    if not root.exists():
        raise ConfigurationError(f"Directory does not exist: {root}")
    if not root.is_dir():
        raise ConfigurationError(f"Not a directory: {root}")

    logger = configure_logging(root)
    started_at = datetime.now().astimezone()
    started_clock = time.perf_counter()
    logger.info("Batch started: root=%s workers=%d", root, args.workers)

    workbook_path = discover_excel(root, args.excel)
    targets = load_target_index(workbook_path)
    for warning in targets.warnings:
        logger.warning(warning)
    logger.info(
        "Loaded %d requested rows (%d Badge keys, %d VIS keys) from %s",
        len(targets.rows),
        len(targets.badge_rows),
        len(targets.vis_rows),
        workbook_path.name,
    )

    validate_runtime()
    logger.info("CodaraScan Panorama runtime initialized")
    plans, discovery_errors, discovered_lots = discover_lots(root, logger)
    stats = RunStats(
        discovered_lots=discovered_lots,
        processed_lots=len(plans),
        processed_pdfs=len(plans),
        planned_pages=sum(plan.pages for plan in plans),
        errors=list(discovery_errors),
    )
    logger.info(
        "Discovered %d lot directories; %d PDFs and %d pages are processable",
        discovered_lots,
        len(plans),
        stats.planned_pages,
    )

    csv_records: list[CsvRecord] = []
    interactive = sys.stdout.isatty()
    noninteractive_step = max(1, stats.planned_pages // 100)
    context = mp.get_context("spawn")
    worker_count = min(args.workers, stats.planned_pages)
    logger.info("Starting %d page worker process(es)", worker_count)
    with context.Pool(processes=worker_count, initializer=initialize_worker) as pool:
        for result in pool.imap_unordered(process_page, iter_tasks(plans), chunksize=1):
            stats.completed_pages += 1
            if result.error:
                stats.failed_pages += 1
                message = (
                    f"Lot {result.lot} | Page {result.page_number}: {result.error}"
                )
                stats.errors.append(message)
                logger.error("%s\n%s", message, result.traceback_text or "")
            else:
                stats.successful_pages += 1
                badge_matches, vis_matches = matched_rows_for_page(result, targets)
                stats.badge_matched_rows.update(badge_matches)
                stats.vis_matched_rows.update(vis_matches)

            # Keep one CSV row for every attempted PDF page. Failed pages have
            # blank identifier fields and remain explicitly documented in the
            # processing log and final report.
            csv_records.extend(csv_records_for_page(result, targets))

            progress = (
                f"Lot {result.lot} | Page {result.page_number}/{result.pages_in_lot} | "
                f"Global {stats.completed_pages}/{stats.planned_pages} pages"
            )
            if interactive:
                print(f"\r{progress:<100}", end="", flush=True)
            elif (
                stats.completed_pages == 1
                or stats.completed_pages == stats.planned_pages
                or stats.completed_pages % noninteractive_step == 0
                or result.error
            ):
                print(progress, flush=True)
    if interactive:
        print()

    output_suffix = workbook_path.suffix
    workbook_output = root / f"{workbook_path.stem}_found{output_suffix}"
    csv_output = root / CSV_NAME
    report_output = root / REPORT_NAME
    found_rows = stats.badge_matched_rows | stats.vis_matched_rows
    write_csv(csv_output, csv_records)
    highlight_workbook(workbook_path, workbook_output, found_rows)

    finished_at = datetime.now().astimezone()
    elapsed = time.perf_counter() - started_clock
    write_report(
        report_output,
        started_at=started_at,
        finished_at=finished_at,
        elapsed=elapsed,
        targets=targets,
        stats=stats,
        workbook_output=workbook_output,
        csv_output=csv_output,
    )
    logger.info(
        "Batch finished in %s: found=%d/%d successful_pages=%d failed_pages=%d",
        format_duration(elapsed),
        len(found_rows),
        len(targets.rows),
        stats.successful_pages,
        stats.failed_pages,
    )
    print(
        f"Finished: found {len(found_rows)}/{len(targets.rows)} requested rows; "
        f"failed pages {stats.failed_pages}. Outputs: {workbook_output.name}, "
        f"{csv_output.name}, {report_output.name}, {LOG_NAME}"
    )
    return 0


def main(argv: Sequence[str] | None = None) -> int:
    try:
        return run(parse_args(argv))
    except ConfigurationError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("Interrupted by user.", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"Fatal error: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    mp.freeze_support()
    raise SystemExit(main())
