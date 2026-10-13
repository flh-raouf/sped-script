#!/usr/bin/env python3
"""Cumulative reports of the working root, shared by the three stages.

Each stage records its page results in one SQLite store. The readable files are
regenerated from that store, so they always cover every lot processed so far,
whichever launcher (``main.py`` or ``run_all.py``) ran the stage:

    Rapports/
      resume.txt                         overall progress
      suivi_lots.xlsx                    one row per lot
      <liste>_found.xlsx                 the search list, every row found so far highlighted
      Pages/pages_lots_0001-0200.xlsx    one row per page, completed by each stage
      Lots/<lot>/                        pages.xlsx, rapport.txt and one CSV per stage
      Executions/                        one report per stage launch and per complete run
      .interne/                          store, OCR cache, logs, timings: read by the scripts only

A file that cannot be replaced (typically because it is open in Excel) stays
marked as stale and is rewritten by the next stage, or by running this script.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import sys
import tempfile
from collections import defaultdict
from collections.abc import Callable, Iterable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

if __package__:
    from . import script as barcode
else:
    import script as barcode

REPORTS = 'Rapports'
INTERNAL = '.interne'
LEGACY_REPORTS = 'Reports'
LOTS_PER_PART = 200   # about 120 000 rows at 600 pages per lot; Excel stops at 1 048 576
EXCEL_ROW_LIMIT = 1_048_575

STAGES = ('barcode', 'ocr', 'review')
STAGE_LABELS = {'barcode': 'Code-barres', 'ocr': 'OCR', 'review': 'Classification finale'}
STAGE_FILES = {'barcode': 'code-barres.csv', 'ocr': 'ocr.csv', 'review': 'classification.csv'}
EXECUTION_LABELS = {'barcode': 'code-barres', 'ocr': 'ocr', 'review': 'classification',
                    'pipeline': 'lancement complet'}

CLASSIFIED, WAITING_OCR, WAITING_REVIEW, TO_REVIEW = 'CLASSIFIED', 'OCR', 'PENDING', 'REVIEW'
STATE_LABELS = {CLASSIFIED: 'Classée', WAITING_OCR: "En attente d'OCR",
                WAITING_REVIEW: 'En attente de classification', TO_REVIEW: 'À revoir'}

PAGE_COLUMNS = ('lot_page', 'BDG', 'VIN', 'VIS', 'SEQ', 'SEQ_9', 'NOF', 'État', 'Classée par',
                'Liste de recherche', 'Code-barres', 'OCR', 'Classification', 'Erreur', 'PDF')
PAGE_WIDTHS = (14, 10, 22, 12, 22, 12, 12, 28, 22, 18, 20, 24, 28, 30, 40)
LOT_COLUMNS = ('Lot', 'Pages', 'Classées', 'Par code-barres', 'Par OCR', 'Par classification',
               "En attente d'OCR", 'En attente de classification', 'À revoir', 'Erreurs',
               'Pages trouvées dans la liste', 'Code-barres le', 'Durée code-barres',
               'Durée OCR (cumulée)', 'Durée classification (cumulée)', 'État')
LOT_WIDTHS = (8, 8, 10, 15, 10, 17, 16, 27, 10, 9, 26, 20, 18, 20, 28, 10)
LOT_STATUS_LABELS = {'SAVED': 'enregistré', 'SAVED_WITH_ISSUES': 'enregistré avec des remarques'}

SCHEMA = """
CREATE TABLE IF NOT EXISTS pages (
    lot TEXT NOT NULL, page INTEGER NOT NULL, lot_number INTEGER NOT NULL,
    state TEXT NOT NULL, method TEXT NOT NULL DEFAULT '',
    bdg TEXT NOT NULL DEFAULT '', vin TEXT NOT NULL DEFAULT '', vis TEXT NOT NULL DEFAULT '',
    seq TEXT NOT NULL DEFAULT '', seq9 TEXT NOT NULL DEFAULT '', nof TEXT NOT NULL DEFAULT '',
    badges TEXT NOT NULL DEFAULT '', decoded_vis TEXT NOT NULL DEFAULT '',
    barcode TEXT NOT NULL DEFAULT '', ocr TEXT NOT NULL DEFAULT '', review TEXT NOT NULL DEFAULT '',
    error TEXT NOT NULL DEFAULT '', pdf TEXT NOT NULL DEFAULT '', updated TEXT NOT NULL,
    PRIMARY KEY (lot, page));
CREATE INDEX IF NOT EXISTS pages_by_number ON pages (lot_number, lot, page);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY, lot TEXT NOT NULL, page INTEGER NOT NULL, stage TEXT NOT NULL,
    run TEXT NOT NULL, at TEXT NOT NULL, seconds REAL NOT NULL, state TEXT NOT NULL,
    error TEXT NOT NULL, detail TEXT NOT NULL);
CREATE INDEX IF NOT EXISTS events_by_lot ON events (lot, stage, page);
CREATE TABLE IF NOT EXISTS lot_runs (
    lot TEXT NOT NULL, stage TEXT NOT NULL, run TEXT NOT NULL, started TEXT NOT NULL,
    finished TEXT NOT NULL, seconds REAL NOT NULL, status TEXT NOT NULL, notes TEXT NOT NULL,
    PRIMARY KEY (lot, stage, run));
CREATE TABLE IF NOT EXISTS stale (kind TEXT NOT NULL, key TEXT NOT NULL, PRIMARY KEY (kind, key));
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""


# --- locations ----------------------------------------------------------------

def reports_dir(root: Path) -> Path:
    return Path(root) / REPORTS


def internal_dir(root: Path) -> Path:
    return reports_dir(root) / INTERNAL


def runs_dir(root: Path, stage: str) -> Path:
    """Technical files of each launch: log, timings, transfer journal."""
    return internal_dir(root) / 'executions' / stage


def legacy_runs_dir(root: Path, stage: str) -> Path:
    return Path(root) / LEGACY_REPORTS / stage


def store_path(root: Path) -> Path:
    return internal_dir(root) / 'rapports.sqlite3'


STATE_FILES = (('cache_ocr.sqlite3', '.ocr_cache.sqlite3'),
               ('review_assignments.jsonl', 'review_assignments.jsonl'),
               ('pipeline_state.json', 'pipeline_state.json'))


def _state_file(root: Path, name: str) -> Path:
    """The state file, at its former location as long as it could not be moved."""
    legacy_name = dict(STATE_FILES)[name]
    new, old = internal_dir(root) / name, Path(root) / LEGACY_REPORTS / legacy_name
    return old if old.exists() and not new.exists() else new


def cache_path(root: Path) -> Path:
    return _state_file(root, 'cache_ocr.sqlite3')


def provenance_path(root: Path) -> Path:
    return _state_file(root, 'review_assignments.jsonl')


def registry_path(root: Path) -> Path:
    return _state_file(root, 'pipeline_state.json')


def prepare(root: Path) -> None:
    """Create ``Rapports/`` and move the state kept in the former ``Reports/`` folder.

    The OCR cache, the Review provenance and the lot registry are state rather
    than reports: once moved, the old ``Reports/`` folder can be deleted safely.
    A file that cannot be moved keeps being used where it is.
    """
    internal = internal_dir(root)
    internal.mkdir(parents=True, exist_ok=True)
    for name, legacy_name in STATE_FILES:
        old, new = Path(root) / LEGACY_REPORTS / legacy_name, internal / name
        if not old.exists() or new.exists():
            continue
        try:
            if old.suffix == '.sqlite3':
                # Fold the write-ahead log into the file so that it moves alone.
                connection = sqlite3.connect(old, timeout=60)
                try:
                    connection.execute('PRAGMA journal_mode=DELETE')
                finally:
                    connection.close()
                if any(Path(f'{old}{suffix}').exists() for suffix in ('-wal', '-journal')):
                    continue
            os.replace(old, new)
        except (OSError, sqlite3.Error):
            continue


def execution_report_path(root: Path, started_at: datetime, kind: str) -> Path:
    directory = reports_dir(root) / 'Executions'
    directory.mkdir(parents=True, exist_ok=True)
    stem = f'{started_at:%Y-%m-%d_%H-%M-%S} - {EXECUTION_LABELS[kind]}'
    path, copy = directory / f'{stem}.txt', 2
    while path.exists():
        path, copy = directory / f'{stem} ({copy}).txt', copy + 1
    path.touch()
    return path


def relative(root: Path, path: Path | str | None) -> str:
    if not path:
        return ''
    path = Path(path)
    try:
        return path.resolve().relative_to(Path(root).resolve()).as_posix()
    except ValueError:
        return str(path)


# --- store ----------------------------------------------------------------------

def page_identity(name: str) -> tuple[str, int] | None:
    """(lot, page) of a one-page PDF named ``<lot>_<page>``."""
    match = re.fullmatch(r'([0-9]+)_([0-9]+)', name)
    return (match[1], int(match[2])) if match else None


def part_of(lot_number: int) -> int:
    return (max(lot_number, 1) - 1) // LOTS_PER_PART


def part_name(part: int) -> str:
    first = part * LOTS_PER_PART + 1
    last = first + LOTS_PER_PART - 1
    width = max(4, len(str(last)))
    return f'pages_lots_{first:0{width}d}-{last:0{width}d}.xlsx'


@dataclass(frozen=True)
class PageEntry:
    """What one stage concluded for one page."""
    lot: str
    page: int
    state: str                        # CLASSIFIED, OCR, PENDING or REVIEW
    row: tuple[str, ...]              # lot_page, BDG, VIN, VIS, SEQ, SEQ_9, NOF
    status: str                       # the stage's own decision code
    detail: dict[str, str] = field(default_factory=dict)   # the stage's CSV columns
    pdf: str = ''
    error: str = ''
    badges: tuple[str, ...] = ()      # identifiers read on the page, to match the search list
    decoded_vis: tuple[str, ...] = ()
    seconds: float = 0.0


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec='seconds')


def _merge(saved: str, values: Iterable[str]) -> str:
    return ';'.join(sorted({value for value in saved.split(';') if value} | {value for value in values if value}))


class Store:
    """Every page result recorded by the stages; the single source of the reports."""

    def __init__(self, root: Path):
        self.root = Path(root)
        path = store_path(self.root)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path, timeout=60, isolation_level=None)
        self.connection.execute('PRAGMA journal_mode=WAL')
        self.connection.execute('PRAGMA synchronous=FULL')
        self.connection.executescript(SCHEMA)

    def close(self) -> None:
        self.connection.close()

    def __enter__(self) -> Store:
        return self

    def __exit__(self, *exc) -> None:
        self.close()

    @contextmanager
    def _transaction(self) -> Iterator[None]:
        self.connection.execute('BEGIN IMMEDIATE')
        try:
            yield
        except BaseException:
            self.connection.execute('ROLLBACK')
            raise
        self.connection.execute('COMMIT')

    # settings

    def set_search_list(self, path: Path) -> None:
        self.connection.execute('INSERT OR REPLACE INTO settings VALUES (?, ?)',
                                ('search_list', str(Path(path).resolve())))

    def search_list(self) -> Path | None:
        row = self.connection.execute("SELECT value FROM settings WHERE key='search_list'").fetchone()
        return Path(row[0]) if row else None

    # writing

    def _mark(self, lot: str) -> None:
        self.connection.executemany('INSERT OR IGNORE INTO stale VALUES (?, ?)',
                                    (('lot', lot), ('part', str(part_of(int(lot)))), ('global', '')))

    def _write_page(self, stage: str, entry: PageEntry, saved: tuple[str, str] | None) -> None:
        badges = _merge(saved[0] if saved else '', entry.badges)
        decoded = _merge(saved[1] if saved else '', entry.decoded_vis)
        if entry.error and saved:
            # A failed page keeps what earlier stages established about it.
            self.connection.execute(
                f'UPDATE pages SET state=?, {stage}=?, error=?, pdf=?, updated=? WHERE lot=? AND page=?',
                (entry.state, entry.status, entry.error, entry.pdf, now(), entry.lot, entry.page))
            return
        row = tuple(entry.row[1:7]) + ('',) * (6 - len(entry.row[1:7]))
        method = stage if entry.state == CLASSIFIED else ''
        self.connection.execute(
            f'INSERT INTO pages (lot, page, lot_number, state, method, bdg, vin, vis, seq, seq9, nof, '
            f'badges, decoded_vis, {stage}, error, pdf, updated) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) '
            f'ON CONFLICT (lot, page) DO UPDATE SET state=excluded.state, method=excluded.method, '
            f'bdg=excluded.bdg, vin=excluded.vin, vis=excluded.vis, seq=excluded.seq, seq9=excluded.seq9, '
            f'nof=excluded.nof, badges=excluded.badges, decoded_vis=excluded.decoded_vis, '
            f'{stage}=excluded.{stage}, error=excluded.error, pdf=excluded.pdf, updated=excluded.updated',
            (entry.lot, entry.page, int(entry.lot), entry.state, method, *row, badges, decoded,
             entry.status, entry.error, entry.pdf, now()))

    def _write_event(self, stage: str, run: str, entry: PageEntry, replaces: int | None) -> int:
        values = (now(), entry.seconds, entry.state, entry.error, json.dumps(entry.detail, ensure_ascii=False))
        if replaces is not None:
            self.connection.execute('UPDATE events SET at=?, seconds=?, state=?, error=?, detail=? WHERE id=?',
                                    (*values, replaces))
            return replaces
        cursor = self.connection.execute(
            'INSERT INTO events (lot, page, stage, run, at, seconds, state, error, detail) VALUES (?,?,?,?,?,?,?,?,?)',
            (entry.lot, entry.page, stage, run, *values))
        return cursor.lastrowid

    @contextmanager
    def replace_lot(self, lot: str, run: str, entries: list[PageEntry]) -> Iterator[None]:
        """Record a scanned lot in place of any earlier scan of it.

        Committed when the block completes, i.e. once the lot is published;
        nothing is kept if the publication fails.
        """
        with self._transaction():
            for table in ('pages', 'events', 'lot_runs'):
                self.connection.execute(f'DELETE FROM {table} WHERE lot=?', (lot,))
            for entry in entries:
                self._write_page('barcode', entry, None)
                self._write_event('barcode', run, entry, None)
            self._mark(lot)
            yield

    def record_lot_run(self, lot: str, stage: str, run: str, *, started: datetime, finished: datetime,
                       seconds: float, status: str, notes: Iterable[str] = ()) -> None:
        with self._transaction():
            self.connection.execute(
                'INSERT OR REPLACE INTO lot_runs VALUES (?,?,?,?,?,?,?,?)',
                (lot, stage, run, started.isoformat(timespec='seconds'), finished.isoformat(timespec='seconds'),
                 seconds, status, '\n'.join(notes)))
            self._mark(lot)

    def record(self, stage: str, run: str, entry: PageEntry, *, replaces: int | None = None) -> int:
        """Record one page; ``replaces`` corrects the page's earlier record of this run.

        OCR and final classification record a page *before* moving its PDF: if the
        move fails or is interrupted, the page is still in the stage's input
        folder, and is recorded again when it is processed again.
        """
        if stage not in STAGES:
            raise ValueError(f'Unknown stage: {stage}')
        with self._transaction():
            saved = self.connection.execute('SELECT badges, decoded_vis FROM pages WHERE lot=? AND page=?',
                                            (entry.lot, entry.page)).fetchone()
            self._write_page(stage, entry, saved)
            event = self._write_event(stage, run, entry, replaces)
            self._mark(entry.lot)
        return event

    # reading

    def metadata(self) -> dict[str, tuple[str, ...]]:
        """lot_page -> (lot_page, BDG, VIN, VIS, SEQ, SEQ_9, NOF) of every classified page."""
        rows = self.connection.execute(
            "SELECT lot, page, bdg, vin, vis, seq, seq9, nof FROM pages WHERE state=? AND vis<>''", (CLASSIFIED,))
        return {f'{lot}_{page}': (f'{lot}_{page}', *values) for lot, page, *values in rows}

    def stale(self, kind: str) -> list[str]:
        return [row[0] for row in self.connection.execute('SELECT key FROM stale WHERE kind=?', (kind,))]

    def clear(self, kind: str, key: str) -> None:
        self.connection.execute('DELETE FROM stale WHERE kind=? AND key=?', (kind, key))

    def mark_everything(self) -> None:
        with self._transaction():
            for (lot,) in self.connection.execute('SELECT DISTINCT lot FROM pages').fetchall():
                self._mark(lot)
            self.connection.execute("INSERT OR IGNORE INTO stale VALUES ('global', '')")

    def pages(self, *, lot: str | None = None, part: int | None = None) -> Iterator[tuple]:
        columns = ('SELECT lot, page, bdg, vin, vis, seq, seq9, nof, state, method, badges, decoded_vis, '
                   'barcode, ocr, review, error, pdf FROM pages ')
        if lot is not None:
            yield from self.connection.execute(columns + 'WHERE lot=? ORDER BY page', (lot,))
        else:
            first = part * LOTS_PER_PART + 1
            low = 0 if part == 0 else first
            yield from self.connection.execute(
                columns + 'WHERE lot_number BETWEEN ? AND ? ORDER BY lot_number, lot, page',
                (low, first + LOTS_PER_PART - 1))

    def identifiers(self) -> Iterator[tuple[str, str, str, str]]:
        yield from self.connection.execute('SELECT lot, badges, decoded_vis, vis FROM pages')

    def events(self, lot: str, stage: str) -> list[tuple]:
        return self.connection.execute(
            'SELECT run, at, seconds, state, error, detail FROM events WHERE lot=? AND stage=? ORDER BY page, id',
            (lot, stage)).fetchall()

    def lot_runs(self, lot: str) -> list[tuple]:
        return self.connection.execute(
            'SELECT stage, run, started, finished, seconds, status, notes FROM lot_runs WHERE lot=? ORDER BY run',
            (lot,)).fetchall()

    def lot_counts(self, lot: str | None = None) -> list[tuple]:
        query = ("SELECT lot, COUNT(*), SUM(state='CLASSIFIED'), SUM(method='barcode'), SUM(method='ocr'), "
                 "SUM(method='review'), SUM(state='OCR'), SUM(state='PENDING'), SUM(state='REVIEW'), "
                 "SUM(error<>'') FROM pages ")
        if lot is not None:
            return self.connection.execute(query + 'WHERE lot=? GROUP BY lot', (lot,)).fetchall()
        return self.connection.execute(query + 'GROUP BY lot ORDER BY MIN(lot_number), lot').fetchall()

    def stage_seconds(self) -> dict[tuple[str, str], float]:
        rows = self.connection.execute('SELECT lot, stage, SUM(seconds) FROM events GROUP BY lot, stage')
        return {(lot, stage): seconds for lot, stage, seconds in rows}

    def barcode_runs(self) -> dict[str, tuple[str, float]]:
        rows = self.connection.execute("SELECT lot, started, seconds FROM lot_runs WHERE stage='barcode' ORDER BY run")
        return {lot: (started, seconds) for lot, started, seconds in rows}


# --- search list -----------------------------------------------------------------

class Matcher:
    """Which rows of the search list a page's identifiers designate (as the barcode stage does)."""

    def __init__(self, targets: barcode.TargetIndex | None):
        self.targets = targets

    def rows(self, badges: str, decoded_vis: str, vis: str) -> set[barcode.RowRef]:
        if self.targets is None:
            return set()
        found: set[barcode.RowRef] = set()
        for badge in filter(None, badges.split(';')):
            found.update(self.targets.badge_rows.get(badge, ()))
        for value in {*filter(None, decoded_vis.split(';')), vis} - {''}:
            found.update(self.targets.vis_rows.get(value, ()))
        return found


def load_targets(path: Path | None, warnings: list[str]) -> barcode.TargetIndex | None:
    if path is None:
        warnings.append('Aucune liste de recherche connue : la liste coloriée n\'est pas mise à jour.')
        return None
    if not path.is_file():
        warnings.append(f'Liste de recherche introuvable : {path}')
        return None
    try:
        return barcode.load_target_index(path)
    except Exception as exc:  # noqa: BLE001 -- the reports must not stop on a damaged list
        warnings.append(f'Liste de recherche illisible ({path.name}) : {exc}')
        return None


# --- files ------------------------------------------------------------------------

def _clean(value):
    if value is None or value == '':
        return None
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def write_workbook(path: Path, title: str, columns: tuple[str, ...], widths: tuple[int, ...],
                   rows: Iterable[tuple]) -> int:
    """Write a filterable sheet, then replace ``path`` in one step; return the row count."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header = []
    for label in columns:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='253B53')
        header.append(cell)
    sheet.append(header)
    count = 0
    for row in rows:
        count += 1
        if count > EXCEL_ROW_LIMIT:
            raise ValueError(f'{path.name} dépasse la limite de lignes d\'Excel')
        sheet.append([_clean(value) for value in row])
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{count + 1}'
    descriptor, name = tempfile.mkstemp(dir=path.parent, prefix=f'.{path.stem}.', suffix=path.suffix)
    os.close(descriptor)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return count


def write_csv(path: Path, columns: Iterable[str], rows: Iterable[Iterable]) -> None:
    import csv

    def write(handle):
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows)
    barcode.atomic_text_write(path, write, newline='', encoding='utf-8-sig')


def write_text(path: Path, text: str) -> None:
    barcode.atomic_text_write(path, lambda handle: handle.write(text), encoding='utf-8')


def number(value: int | float | None) -> str:
    return f'{int(value or 0):,}'.replace(',', ' ')


def moment(value: str | None) -> str:
    if not value:
        return ''
    try:
        return datetime.fromisoformat(value).strftime('%d/%m/%Y %H:%M:%S')
    except ValueError:
        return value


def run_moment(run: str) -> str:
    try:
        return datetime.strptime(run, '%Y%m%d_%H%M%S_%f').strftime('%d/%m/%Y %H:%M:%S')
    except ValueError:
        return run


def page_row(record: tuple, found: bool) -> tuple:
    (lot, page, bdg, vin, vis, seq, seq9, nof, state, method, _badges, _decoded,
     barcode_status, ocr_status, review_status, error, pdf) = record
    return (f'{lot}_{page}', bdg, vin, vis, seq, seq9, nof, STATE_LABELS.get(state, state),
            STAGE_LABELS.get(method, ''), 'Trouvée' if found else '', barcode_status, ocr_status,
            review_status, error, pdf)


def page_rows(records: Iterable[tuple], matcher: Matcher) -> Iterator[tuple]:
    for record in records:
        yield page_row(record, bool(matcher.rows(record[10], record[11], record[4])))


def stage_rows(events: list[tuple]) -> tuple[list[str], list[list]]:
    """The stage's own columns, preceded by when the page was recorded."""
    details = [json.loads(event[5]) for event in events]
    columns = list(dict.fromkeys(key for detail in details for key in detail))
    rows = [[run_moment(event[0]), moment(event[1]), *(detail.get(key, '') for key in columns)]
            for event, detail in zip(events, details)]
    return ['Exécution', 'Enregistrée le', *columns], rows


def lot_report(lot: str, counts: tuple | None, runs: list[tuple], stage_events: dict[str, list[tuple]]) -> str:
    lines = [f'Lot {lot} — rapport mis à jour le {datetime.now():%d/%m/%Y %H:%M:%S}', '']
    if counts:
        _, total, classified, by_barcode, by_ocr, by_review, ocr, pending, review, errors = counts
        lines += [f'Pages : {number(total)}',
                  f'  Classées : {number(classified)} (code-barres : {number(by_barcode)} · OCR : {number(by_ocr)}'
                  f' · classification finale : {number(by_review)})',
                  f'  À revoir : {number(review)}',
                  f"  En attente d'OCR : {number(ocr)}",
                  f'  En attente de classification : {number(pending)}',
                  f'  En erreur : {number(errors)}']
    for stage in STAGES:
        lines += ['', STAGE_LABELS[stage]]
        events = stage_events[stage]
        if not events:
            lines.append('  Pas encore exécuté pour ce lot.')
            continue
        by_run = defaultdict(list)
        for event in events:
            by_run[event[0]].append(event)
        timings = {run: (started, finished, seconds, status, notes)
                   for run_stage, run, started, finished, seconds, status, notes in runs if run_stage == stage}
        for run in sorted(by_run):
            items = by_run[run]
            classified = sum(event[3] == CLASSIFIED and not event[4] for event in items)
            errors = sum(bool(event[4]) for event in items)
            lines.append(f'  Exécution du {run_moment(run)}')
            if run in timings:
                started, finished, seconds, status, notes = timings[run]
                lines.append(f'    Début : {moment(started)} · Fin : {moment(finished)} · '
                             f'Durée : {barcode.format_duration(seconds)}')
            else:
                # OCR and classification take the pages of several lots together.
                lines.append(f'    Pages enregistrées du {moment(min(event[1] for event in items))} au '
                             f'{moment(max(event[1] for event in items))} · Durée cumulée des pages : '
                             f'{barcode.format_duration(sum(event[2] for event in items))}')
                status, notes = '', ''
            onward = {'barcode': "envoyées en OCR", 'ocr': 'envoyées en classification finale',
                      'review': 'à revoir'}[stage]
            text = (f'    {number(len(items))} pages : {number(classified)} classées, '
                    f'{number(len(items) - classified - errors)} {onward}')
            # A page the barcode stage cannot read simply goes to OCR; its remarks list it.
            lines.append(text if stage == 'barcode' else f'{text}, {number(errors)} en erreur')
            if status:
                lines.append(f'    Statut : {LOT_STATUS_LABELS.get(status, status)}')
            if notes:
                lines.append('    Remarques :')
                lines.extend(f'      - {note}' for note in notes.splitlines())
    return '\n'.join(lines) + '\n'


def write_lot(root: Path, store: Store, lot: str, matcher: Matcher) -> None:
    directory = reports_dir(root) / 'Lots' / lot
    counts = store.lot_counts(lot)
    if not counts:
        return
    directory.mkdir(parents=True, exist_ok=True)
    write_workbook(directory / 'pages.xlsx', 'Pages', PAGE_COLUMNS, PAGE_WIDTHS,
                   page_rows(store.pages(lot=lot), matcher))
    stage_events = {stage: store.events(lot, stage) for stage in STAGES}
    for stage, events in stage_events.items():
        path = directory / STAGE_FILES[stage]
        if events:
            write_csv(path, *stage_rows(events))
        else:
            path.unlink(missing_ok=True)
    write_text(directory / 'rapport.txt', lot_report(lot, counts[0], store.lot_runs(lot), stage_events))


def write_part(root: Path, store: Store, part: int, matcher: Matcher) -> None:
    directory = reports_dir(root) / 'Pages'
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / part_name(part)
    rows = list(page_rows(store.pages(part=part), matcher))
    if rows:
        write_workbook(path, 'Pages', PAGE_COLUMNS, PAGE_WIDTHS, rows)
    else:
        path.unlink(missing_ok=True)


def write_globals(root: Path, store: Store, targets: barcode.TargetIndex | None, list_path: Path | None,
                  problems: list[str]) -> None:
    matcher = Matcher(targets)
    found_rows: set[barcode.RowRef] = set()
    found_pages: dict[str, int] = defaultdict(int)
    for lot, badges, decoded, vis in store.identifiers():
        rows = matcher.rows(badges, decoded, vis)
        if rows:
            found_rows |= rows
            found_pages[lot] += 1
    counts = store.lot_counts()
    seconds = store.stage_seconds()
    scans = store.barcode_runs()
    lots = []
    for lot, total, classified, by_barcode, by_ocr, by_review, ocr, pending, review, errors in counts:
        started, scan_seconds = scans.get(lot, ('', None))
        lots.append((lot, total, classified, by_barcode, by_ocr, by_review, ocr, pending, review, errors,
                     found_pages.get(lot, 0) if targets else None, moment(started),
                     barcode.format_duration(scan_seconds) if scan_seconds is not None else '',
                     barcode.format_duration(seconds[lot, 'ocr']) if (lot, 'ocr') in seconds else '',
                     barcode.format_duration(seconds[lot, 'review']) if (lot, 'review') in seconds else '',
                     'En cours' if ocr or pending else 'Terminé'))
    directory = reports_dir(root)
    write_workbook(directory / 'suivi_lots.xlsx', 'Lots', LOT_COLUMNS, LOT_WIDTHS, lots)
    if targets is not None and list_path is not None:
        output = directory / f'{list_path.stem}_found{list_path.suffix}'
        barcode.highlight_workbook(list_path, output, found_rows)
    totals = [sum(row[index] or 0 for row in counts) for index in range(1, 10)]
    total, classified, by_barcode, by_ocr, by_review, ocr, pending, review, errors = totals
    share = f' ({100 * classified / total:.1f} %)'.replace('.', ',') if total else ''
    lines = [f'Rapports SPED — mis à jour le {datetime.now():%d/%m/%Y %H:%M:%S}', '',
             f'Lots : {number(len(lots))} ({number(sum(row[-1] == "Terminé" for row in lots))} terminés, '
             f'{number(sum(row[-1] == "En cours" for row in lots))} en cours)',
             f'Pages : {number(total)}',
             f'  Classées : {number(classified)}{share}',
             f'    par code-barres : {number(by_barcode)}',
             f'    par OCR : {number(by_ocr)}',
             f'    par classification finale : {number(by_review)}',
             f"  En attente d'OCR : {number(ocr)}",
             f'  En attente de classification : {number(pending)}',
             f'  À revoir (dossier Review) : {number(review)}',
             f'  En erreur : {number(errors)}']
    if targets is not None and list_path is not None:
        lines.append(f'Liste de recherche « {list_path.name} » : {number(len(found_rows))} lignes trouvées '
                     f'sur {number(len(targets.rows))} ({list_path.stem}_found{list_path.suffix})')
    lines += ['', 'Fichiers :',
              '  suivi_lots.xlsx — une ligne par lot',
              f'  Pages/ — une ligne par page, {LOTS_PER_PART} lots par fichier',
              '  Lots/<lot>/ — le détail de chaque lot (pages, rapport, une feuille par étape)',
              '  Executions/ — le rapport de chaque lancement']
    if problems:
        lines += ['', 'Non mis à jour (fichier ouvert dans Excel ?), réessayé au prochain traitement :',
                  *(f'  - {problem}' for problem in problems)]
    write_text(directory / 'resume.txt', '\n'.join(lines) + '\n')


def refresh(root: Path, store: Store, *, search_list: Path | None = None, complete: bool = True,
            progress: Callable[[str], None] | None = None) -> list[str]:
    """Rewrite the stale report files; return what could not be written.

    ``complete=False`` rewrites only the lot folders: the barcode stage uses it
    after each lot, and refreshes the cumulative files once at the end.
    """
    warnings: list[str] = []
    list_path = search_list or store.search_list()
    targets = load_targets(list_path, warnings) if complete else None
    matcher = Matcher(targets if complete else _quick_targets(list_path))
    problems: list[str] = []
    lots = sorted(store.stale('lot'), key=lambda name: (int(name), name))
    for index, lot in enumerate(lots, 1):
        if progress and (index % 50 == 0 or index == len(lots)):
            progress(f'Lots : {index}/{len(lots)}')
        try:
            write_lot(root, store, lot, matcher)
            store.clear('lot', lot)
        except (OSError, ValueError) as exc:
            problems.append(f'Lots/{lot} : {exc}')
    if complete:
        for key in sorted(store.stale('part'), key=int):
            if progress:
                progress(f'Pages : {part_name(int(key))}')
            try:
                write_part(root, store, int(key), matcher)
                store.clear('part', key)
            except (OSError, ValueError) as exc:
                problems.append(f'Pages/{part_name(int(key))} : {exc}')
        if store.stale('global'):
            if progress:
                progress('Suivi des lots, liste de recherche et résumé')
            try:
                write_globals(root, store, targets, list_path, problems)
                store.clear('global', '')
            except (OSError, ValueError) as exc:
                problems.append(f'fichiers globaux : {exc}')
    return warnings + problems


_TARGET_CACHE: dict[tuple[str, float], barcode.TargetIndex] = {}


def _quick_targets(path: Path | None) -> barcode.TargetIndex | None:
    """The search list, read once per process for the barcode stage's per-lot refreshes."""
    if path is None or not path.is_file():
        return None
    key = (str(path), path.stat().st_mtime)
    if key not in _TARGET_CACHE:
        try:
            _TARGET_CACHE.clear()
            _TARGET_CACHE[key] = barcode.load_target_index(path)
        except Exception:  # noqa: BLE001 -- the complete refresh reports the problem
            return None
    return _TARGET_CACHE[key]


def refresh_quietly(root: Path, store: Store, logger=None, **options) -> list[str]:
    """A report problem never turns a completed document operation into a failure."""
    try:
        problems = refresh(root, store, **options)
    except Exception as exc:  # noqa: BLE001 -- the store keeps everything for the next refresh
        problems = [f'Rapports non mis à jour : {type(exc).__name__}: {exc}']
    for problem in problems:
        print(f'Rapports : {problem}', flush=True)
        if logger is not None:
            logger.warning('Reports | %s', problem)
    return problems


def pending_updates(root: Path) -> int:
    """How many report files are waiting to be rewritten."""
    if not store_path(root).exists():
        return 0
    with Store(root) as store:
        return sum(len(store.stale(kind)) for kind in ('lot', 'part', 'global'))


# --- command line -----------------------------------------------------------------

def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description='Rewrite the cumulative reports of a working root.')
    parser.add_argument('-d', '--directory', type=Path, required=True, help='Working root (racine)')
    parser.add_argument('--excel', type=Path, help='Search list to highlight (default: the last one used)')
    parser.add_argument('--tout', action='store_true', help='Rewrite every file, not only the stale ones')
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    root = args.directory.expanduser().resolve()
    if not root.is_dir():
        print(f'Dossier introuvable : {root}', file=sys.stderr)
        return 2
    prepare(root)
    with Store(root) as store:
        if args.excel is not None:
            store.set_search_list(args.excel if args.excel.is_absolute() else root / args.excel)
        if args.tout:
            store.mark_everything()
        problems = refresh(root, store, progress=lambda text: print(text, flush=True))
    for problem in problems:
        print(f'Attention : {problem}')
    print(f'Rapports : {reports_dir(root)}')
    return 1 if problems else 0


if __name__ == '__main__':
    raise SystemExit(main())
