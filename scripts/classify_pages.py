#!/usr/bin/env python3
"""Split numeric lots and classify individual pages by VIS using the BDD."""
from __future__ import annotations

import argparse
import csv
import io
import multiprocessing as mp
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import time
from collections import defaultdict
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field, replace
from datetime import datetime
from pathlib import Path

if __package__:
    from . import script as barcode
    from . import split_pages as split
    from . import performance
    from . import reporting
else:
    import script as barcode
    import split_pages as split
    import performance
    import reporting


@dataclass(frozen=True)
class VehicleRecord:
    badge: str = ''
    vin: str = ''
    vis: str = ''
    seq: str = ''
    nof: str = ''
    year: int | None = None


@dataclass
class DatabaseIndex:
    badge_vis: dict[str, frozenset[str]] = field(default_factory=dict)
    records_by_vis: dict[str, list[VehicleRecord]] = field(default_factory=dict)


def text_value(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_database(path: Path) -> DatabaseIndex:
    """Stream the consolidated workbook; never choose a contradictory pair."""
    from openpyxl import load_workbook
    mapping: dict[str, set[str]] = defaultdict(set)
    records: dict[str, list[VehicleRecord]] = defaultdict(list)
    recognized = False
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows()
            columns = None
            for number, row in enumerate(rows, 1):
                headers = [barcode.normalize_header(cell.value) for cell in row]
                badge_col = next((i for i, h in enumerate(headers) if h in barcode.BADGE_HEADERS), None)
                vis_col = next((i for i, h in enumerate(headers) if h in barcode.VIS_HEADERS), None)
                if badge_col is not None and vis_col is not None:
                    columns = badge_col, vis_col, headers
                    recognized = True
                    break
                if number >= 50:
                    break
            if columns is None:
                continue
            bi, vi, headers = columns
            seq_col = next((i for i, h in enumerate(headers) if h in {'seq', 'sem', 'sequence'}), None)
            nof_col = next((i for i, h in enumerate(headers) if h == 'nof'), None)
            year_col = next((i for i, h in enumerate(headers) if h in {'annee', 'year'}), None)
            vin_cols = [i for i, h in enumerate(headers) if h in barcode.VIN_HEADERS | {'id1', 'id2', 'id3'}]
            for row in rows:
                badge_cell = row[bi]
                value = badge_cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    if int(value) != value:
                        badge = ''
                    else:
                        badge = str(int(value))
                    if badge and '00000' in badge_cell.number_format and len(badge) < 5:
                        badge = badge.zfill(5)
                else:
                    badge = text_value(value)
                if not re.fullmatch(r'[0-9]{1,6}', badge):
                    badge = ''
                vis = barcode.normalize_vis(row[vi].value)
                if not vis:
                    continue
                if badge:
                    mapping[badge].add(vis)
                vins = {vin for i in vin_cols if (vin := barcode.normalize_vin(row[i].value))
                        and vin[-8:] == vis}
                year = text_value(row[year_col].value) if year_col is not None else ''
                records[vis].append(VehicleRecord(
                    badge=badge, vin=next(iter(vins)) if len(vins) == 1 else '', vis=vis,
                    seq=text_value(row[seq_col].value) if seq_col is not None else '',
                    nof=text_value(row[nof_col].value) if nof_col is not None else '',
                    year=int(year) if re.fullmatch(r'[0-9]{4}', year) else None,
                ))
    finally:
        workbook.close()
    if not recognized or not mapping:
        raise barcode.ConfigurationError('BDD must contain usable BDG/Badge and VIS columns')
    return DatabaseIndex({badge: frozenset(values) for badge, values in mapping.items()}, dict(records))


def page_excel_row(result: barcode.PageResult, vis: str | None, database: DatabaseIndex) -> tuple[str, ...]:
    """Enrich an identified page without guessing absent or conflicting fields."""
    page_id = f'{result.lot}_{result.page_number}'
    if not vis:
        return (page_id, '', '', '', '', '', '')
    candidates = database.records_by_vis.get(vis, [])
    if result.badges:
        candidates = [record for record in candidates if record.badge in result.badges]
    if result.vins:
        exact = [record for record in candidates if record.vin in result.vins]
        if exact:
            candidates = exact
    def unique(values):
        values = {value for value in values if value}
        return next(iter(values)) if len(values) == 1 else ''
    badge = unique(result.badges) or unique(record.badge for record in candidates)
    vin = unique(result.vins) or unique(record.vin for record in candidates)
    seq = unique(record.seq for record in candidates)
    nof = unique(record.nof for record in candidates)
    seq_digits = ''.join(re.findall(r'[0-9]', seq))
    return (page_id, badge, vin, vis, seq, seq_digits[-9:], nof)


def write_page_workbook(path: Path, rows: Iterable[tuple[str, ...]]) -> None:
    """Export one row per page using the application's existing Excel library."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Pages')
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate((20, 14, 24, 16, 30, 18, 18), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header = []
    for label in ('lot_page', 'BDG', 'VIN', 'VIS', 'SEQ', 'SEQ_9', 'NOF'):
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='253B53')
        header.append(cell)
    sheet.append(header)
    row_count = 0
    for row in rows:
        row_count += 1
        if row_count > 1_048_575:
            raise barcode.ConfigurationError('Page report exceeds the Excel worksheet row limit')
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value or None)
            cell.data_type = 's'
            cell.number_format = '@'
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f'A1:G{row_count + 1}'
    workbook.save(path)
    workbook.close()


def resolve_vis(result: barcode.PageResult, database: dict[str, frozenset[str]]) -> tuple[str | None, str]:
    if result.error:
        return None, 'SCAN_ERROR'
    direct = {vin[-8:] for vin in result.vins}
    candidates = set(direct)
    for badge in result.badges:
        mapped = database.get(badge, frozenset())
        if len(mapped) > 1:
            return None, 'BDD_CONFLICT'
        candidates.update(mapped)
    if len(candidates) > 1:
        return None, 'IDENTIFIER_CONFLICT'
    if not candidates:
        return None, 'UNRESOLVED_BADGE' if result.badges else 'NO_BARCODE'
    return next(iter(candidates)), 'BARCODE' if direct else 'BDD_BADGE'


def scan_single_page(task: barcode.PageTask) -> barcode.PageResult:
    # Scan the split PDF itself; restore the original lot/page coordinates.
    try:
        result = barcode.process_page(replace(task, page_index=0))
    finally:
        # Release the file before the parent moves it (also required on Windows).
        barcode._close_worker_document()
    return replace(result, page_number=task.page_index + 1)


def found_rows(result: barcode.PageResult, vis: str | None, targets: barcode.TargetIndex) -> set[barcode.RowRef]:
    badges, vehicles = barcode.matched_rows_for_page(result, targets)
    if vis:
        vehicles.update(targets.vis_rows.get(vis, ()))
    return badges | vehicles


def discover_database(root: Path, explicit: Path | None = None) -> Path:
    if explicit is not None:
        path = explicit.expanduser()
        if not path.is_absolute():
            path = root / path
        path = path.resolve()
        if not path.is_file() or path.suffix.lower() != '.xlsx':
            raise barcode.ConfigurationError(f'BDD must be an existing .xlsx file: {path}')
        return path
    candidates = sorted(path.resolve() for path in root.iterdir()
                        if path.is_file() and path.suffix.lower() == '.xlsx'
                        and path.stem.casefold().startswith('bdd')
                        and not path.stem.casefold().endswith('_found')
                        and not path.name.startswith(('~$', '.')))
    if not candidates:
        raise barcode.ConfigurationError(
            f'No BDD workbook found in {root}. Put BDD_2024_2025_2026.xlsx '
            'directly in the parent folder, or supply --database.')
    if len(candidates) > 1:
        raise barcode.ConfigurationError(
            'Multiple BDD workbooks found (' + ', '.join(path.name for path in candidates)
            + '); select one with --database.')
    return candidates[0]


def discover_search_list(root: Path, database_path: Path, explicit: Path | None = None) -> Path:
    """The search list: given explicitly, or the only other workbook directly in the root."""
    if explicit is None:
        candidates = sorted(path for path in root.iterdir()
                            if path.is_file() and path.suffix.lower() in {'.xlsx', '.xlsm'}
                            and not path.name.startswith(('~$', '.'))
                            and not path.stem.lower().endswith('_found')
                            and path.resolve() != database_path.resolve())
        if len(candidates) != 1:
            raise barcode.ConfigurationError('Select the search list with --excel')
        explicit = candidates[0]
    path = barcode.discover_excel(root, explicit)
    if path.resolve() == database_path.resolve():
        raise barcode.ConfigurationError('Select the search list with --excel, separately from --database')
    return path


CSV_HEADERS = ('Lot', 'Page', 'Badge', 'VIN', 'Decoded VIS', 'VIS', 'Found', 'Status', 'PDF', 'Error')
PAGE_HEADERS = ('lot_page', 'BDG', 'VIN', 'VIS', 'SEQ', 'SEQ_9', 'NOF')


@dataclass
class LotRun:
    name: str
    started_at: datetime
    started_clock: float
    finished_at: datetime | None = None
    elapsed: float = 0.0
    status: str = 'RUNNING'
    planned_pages: int = 0
    processed_pages: int = 0
    classified: int = 0
    ocr: int = 0
    published: bool = False
    issues: list[str] = field(default_factory=list)
    work: Path | None = None
    input_bytes: int = 0
    scan_errors: int = 0


def save_csv(path: Path, headers: Iterable[str], rows: Iterable[Iterable[object]]) -> None:
    def write(handle):
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
    barcode.atomic_text_write(path, write, newline='', encoding='utf-8-sig')


def append_lot_csv(path: Path, rows: list[tuple]) -> None:
    buffer = io.StringIO(newline='')
    csv.writer(buffer).writerows(rows)
    with path.open('ab') as handle:
        offset = handle.tell()
        try:
            handle.write(buffer.getvalue().encode('utf-8'))
            handle.flush()
            os.fsync(handle.fileno())
        except BaseException:
            # Preserve the previously saved CSV prefix if an append is interrupted.
            handle.seek(offset)
            handle.truncate()
            raise


def process_lot(lot: Path, state: LotRun, pool, database: DatabaseIndex,
                targets: barcode.TargetIndex, logger, root=None):
    """Prepare one lot privately; nothing reaches VIS/OCR folders during scanning.

    Returns the CSV records, the page rows, the search-list rows found and the
    pages to record in the cumulative reports.
    """
    from pypdf import PdfReader
    pdf, error = split.discover_lot_pdf(lot)
    if error:
        raise ValueError(error)
    state.input_bytes = pdf.stat().st_size
    work = state.work
    pending = work / '_pending'
    pending.mkdir()
    reader = None
    tasks = []
    try:
        reader = PdfReader(str(pdf), strict=False)
        state.planned_pages = len(reader.pages)
        if not state.planned_pages:
            raise ValueError('PDF contains no pages')
        logger.info('Lot %s | SPLIT START | source=%s | pages=%d', lot.name, pdf.name, state.planned_pages)
        for index, page in enumerate(reader.pages):
            path = pending / f'{lot.name}_{index + 1}.pdf'
            split.write_one_page_pdf(source_page=page, output_path=path)
            tasks.append(barcode.PageTask(lot.name, str(path), index, state.planned_pages, True))
    finally:
        if reader is not None:
            split.close_pdf_reader(reader)
    logger.info('Lot %s | SCAN START | pages=%d', lot.name, state.planned_pages)
    records, page_rows, matches, entries = [], [], set(), []
    for result in pool.imap_unordered(scan_single_page, tasks, chunksize=1):
        state.scan_errors += bool(result.error)
        vis, status = resolve_vis(result, database.badge_vis)
        name = f'{result.lot}_{result.page_number}.pdf'
        destination = work / (vis or 'OCR')
        destination.mkdir(exist_ok=True)
        (pending / name).rename(destination / name)
        state.processed_pages += 1
        state.classified += bool(vis)
        state.ocr += not bool(vis)
        page_matches = found_rows(result, vis, targets)
        matches.update(page_matches)
        record = barcode.csv_records_for_page(result, targets)[0]
        decoded_vis = sorted({vin[-8:] for vin in result.vins})
        pdf_location = (str((root/'Output'/vis if vis else root/'OCR')/name)
                        if root else f'{vis or "OCR"}/{name}')
        records.append((result.lot, result.page_number, record.badge, record.vin,
                        ';'.join(decoded_vis), vis or '', 'FOUND' if page_matches else '',
                        status, pdf_location, result.error or ''))
        page_rows.append((result.page_number, page_excel_row(result, vis, database)))
        entries.append(reporting.PageEntry(
            result.lot, result.page_number, reporting.CLASSIFIED if vis else reporting.WAITING_OCR,
            page_rows[-1][1], status, dict(zip(CSV_HEADERS, map(str, records[-1]))),
            pdf=f'{"Output/" + vis if vis else "OCR"}/{name}', badges=tuple(result.badges),
            decoded_vis=tuple(decoded_vis)))
        if status.endswith('CONFLICT') or result.error:
            message = (f'Lot {result.lot} page {result.page_number}: {status}; '
                       f'Badge={record.badge or "-"}; VIN={record.vin or "-"}; '
                       f'BDD VIS={dict((badge, sorted(database.badge_vis.get(badge, ()))) for badge in result.badges)}; '
                       f'error={result.error or "-"}')
            state.issues.append(message)
            logger.error('%s\n%s', message, result.traceback_text or '')
        print(f'Lot {lot.name} | Page {result.page_number}/{state.planned_pages} | {status}', flush=True)
    if state.processed_pages != state.planned_pages:
        raise RuntimeError(f'Only {state.processed_pages}/{state.planned_pages} page results returned')
    pending.rmdir()
    records.sort(key=lambda row: row[1])
    page_rows.sort(key=lambda row: row[0])
    entries.sort(key=lambda entry: entry.page)
    return records, [row[1] for row in page_rows], matches, entries


def save_lot_reports(work: Path, records: list[tuple], page_rows: list[tuple[str, ...]]) -> None:
    """Old layout: the lot's metadata travels with it, before any PDF is published."""
    reports = work / '_reports'
    reports.mkdir()
    save_csv(reports / 'results.csv', CSV_HEADERS, records)
    save_csv(reports / 'pages.csv', PAGE_HEADERS, page_rows)
    write_page_workbook(reports / 'pages.xlsx', page_rows)


def publish_lot(work: Path, output: Path, lot_name: str, *, vis_output=None, ocr_output=None) -> None:
    """Publish this lot at its boundary; never replace earlier lots' files."""
    def remove_generated_metadata(directory: Path) -> None:
        # macOS creates these sidecars on exFAT drives. This is private staging,
        # so discard its filesystem metadata while preserving every real PDF.
        for entry in directory.iterdir():
            if entry.is_file() and (entry.name.startswith('._') or entry.name == '.DS_Store'):
                entry.unlink(missing_ok=True)

    remove_generated_metadata(work)
    if vis_output is not None:
        # Check every collision before publishing any part of a new-layout lot.
        for directory in work.iterdir():
            if directory.name == '_reports':
                continue
            destination = ocr_output if directory.name == 'OCR' else vis_output/directory.name
            for pdf in directory.glob('*.pdf'):
                if not pdf.name.startswith('.') and (destination/pdf.name).exists():
                    raise FileExistsError(f'Output collision: {destination/pdf.name}')
    for directory in sorted(work.iterdir()):
        if directory.name == '_reports':
            continue
        destination = ((ocr_output if directory.name == 'OCR' else vis_output/directory.name)
                       if vis_output is not None else output/directory.name)
        destination.mkdir(exist_ok=True)
        for pdf in sorted(directory.glob('*.pdf')):
            if pdf.name.startswith('.'):
                continue
            target = destination / pdf.name
            if target.exists():
                raise FileExistsError(f'Output collision: {target}')
            pdf.rename(target)
        remove_generated_metadata(directory)
        directory.rmdir()
    # The complete per-lot metadata is already on disk before any PDF moves.
    if (work / '_reports').exists():
        (work / '_reports').rename(output / 'lot_reports' / lot_name)
    remove_generated_metadata(work)
    work.rmdir()


def saved_page_rows(output: Path, states: list[LotRun]) -> Iterator[tuple[str, ...]]:
    for state in states:
        if not state.published:
            continue
        with (output / 'lot_reports' / state.name / 'pages.csv').open(encoding='utf-8-sig', newline='') as handle:
            rows = csv.reader(handle)
            next(rows)
            yield from (tuple(row) for row in rows)


def finish_lot(state: LotRun, output: Path, logger, lot_report: bool = True) -> None:
    state.finished_at = datetime.now().astimezone()
    state.elapsed = time.perf_counter() - state.started_clock
    duration = barcode.format_duration(state.elapsed)
    summary = (f'Lot {state.name} | {state.status} | duration={duration} ({state.elapsed:.3f} seconds) | '
               f'pages={state.processed_pages}/{state.planned_pages} | VIS={state.classified} | OCR={state.ocr} | '
               f'published={state.published}')
    logger.info('%s | finish=%s', summary, state.finished_at.isoformat(timespec='seconds'))
    print(summary, flush=True)
    if not lot_report:
        return
    report_dir = output / 'lot_reports' / state.name
    report_dir.mkdir(exist_ok=True)
    lines = [summary, f'Start time: {state.started_at.isoformat(timespec="seconds")}',
             f'Finish time: {state.finished_at.isoformat(timespec="seconds")}',
             f'Duration seconds: {state.elapsed:.3f}', *state.issues]
    if state.work and state.work.exists():
        lines.append(f'Unfinished lot files retained in: {state.work}')
    barcode.atomic_text_write(report_dir / 'report.txt', lambda handle: handle.write('\n'.join(lines) + '\n'))


def write_batch_report(output: Path, *, started_at: datetime, elapsed: float, status: str,
                       database_path: Path, targets: barcode.TargetIndex, states: list[LotRun],
                       lots_discovered: int, matches: set, issues: list[str], finished_at=None,
                       report_path: Path | None = None) -> None:
    """``report_path`` is the execution report of the new layout; the old one keeps ``report.txt``."""
    published = [state for state in states if state.published]
    if report_path is None:
        files = ['Per-lot page workbooks: lot_reports/<lot>/pages.xlsx',
                 f'Combined page workbook: {"pages.xlsx" if (output / "pages.xlsx").exists() else "pending"}']
    else:
        files = [f'Reports: {report_path.parent.parent}', f'Technical files: {output}']
    lines = [f'Batch status: {status}', f'Start time: {started_at.isoformat(timespec="seconds")}',
             f'Finish time: {finished_at.isoformat(timespec="seconds") if finished_at else ""}',
             f'Total processing time: {barcode.format_duration(elapsed)}',
             f'Total processing seconds: {elapsed:.3f}', '', f'BDD: {database_path.name}',
             f'Search list: {targets.workbook_path.name}', *files,
             f'Lots discovered: {lots_discovered}', f'Lots published: {len(published)}',
             f'Pages processed: {sum(state.processed_pages for state in published)}',
             f'Pages classified by VIS: {sum(state.classified for state in published)}',
             f'Pages in OCR: {sum(state.ocr for state in published)}',
             f'Found: {len(matches)}/{len(targets.rows)}', '', 'Lot durations:']
    lines.extend(f'- Lot {s.name}: {s.status}; {barcode.format_duration(s.elapsed)} '
                 f'({s.elapsed:.3f} seconds); pages {s.processed_pages}/{s.planned_pages}' for s in states)
    lines.extend(['', f'Issues: {len(issues)}', *[f'- {issue}' for issue in issues]])
    barcode.atomic_text_write(report_path or output / 'report.txt',
                              lambda handle: handle.write('\n'.join(lines) + '\n'))
    save_csv(output / 'lots.csv',
             ('Lot', 'Status', 'Start', 'Finish', 'Duration', 'Seconds', 'Planned pages', 'Processed pages', 'VIS', 'OCR', 'Published'),
             ((s.name, s.status, s.started_at.isoformat(timespec='seconds'),
               s.finished_at.isoformat(timespec='seconds') if s.finished_at else '',
               barcode.format_duration(s.elapsed), f'{s.elapsed:.3f}', s.planned_pages,
               s.processed_pages, s.classified, s.ocr, s.published) for s in states))


def run(args: argparse.Namespace) -> int:
    started_clock = time.perf_counter()
    started_at = datetime.now().astimezone()
    print(f'Batch START | {started_at.isoformat(timespec="seconds")} | Loading BDD and scanner', flush=True)
    root = split.validate_root(args.directory)
    database_path = discover_database(root, args.database)
    target_path = discover_search_list(root, database_path, args.excel)
    targets = barcode.load_target_index(target_path)
    database = load_database(database_path)
    barcode.validate_runtime()
    lots = split.discover_lot_directories(root)
    if not lots:
        raise barcode.ConfigurationError('No numeric Lot folders found')
    run_id = started_at.strftime('%Y%m%d_%H%M%S_%f')
    report_path = None
    if args.legacy_layout:
        output = root / 'output'
        initial = Path(tempfile.mkdtemp(prefix='.classify_pages_', dir=root))
        try:
            split.publish_staging_directory(initial, output, split.unique_backup_path(output))
        finally:
            if initial.exists():
                shutil.rmtree(initial)
    else:
        reporting.prepare(root)
        output = reporting.runs_dir(root, 'barcode') / run_id
        output.mkdir(parents=True)
        (root/'Output').mkdir(exist_ok=True)
        (root/'OCR').mkdir(exist_ok=True)
        report_path = reporting.execution_report_path(root, started_at, 'barcode')
    logger = barcode.configure_logging(output)
    logger.info('Batch started: %s | root=%s | workers=%d | lots=%d',
                started_at.isoformat(timespec='seconds'), root, args.workers, len(lots))
    logger.info('Setup complete | BDD badges=%d | requested rows=%d', len(database.badge_vis), len(targets.rows))
    if args.legacy_layout:
        (output / 'OCR').mkdir()
        (output / 'lot_reports').mkdir()
    incomplete = output / '_incomplete'
    incomplete.mkdir()
    matches: set[barcode.RowRef] = set()
    states: list[LotRun] = []
    issues = list(targets.warnings)
    active = store = None
    interrupted = False
    def update_report(status='RUNNING', finished_at=None, elapsed=None):
        write_batch_report(output, started_at=started_at,
                           elapsed=time.perf_counter() - started_clock if elapsed is None else elapsed,
                           status=status, database_path=database_path, targets=targets, states=states,
                           lots_discovered=len(lots), matches=matches, issues=issues, finished_at=finished_at,
                           report_path=report_path)
    try:
        if args.legacy_layout:
            save_csv(output / 'results.csv', CSV_HEADERS, [])
        else:
            store = reporting.Store(root)
            store.set_search_list(target_path)
        update_report()
        setup_seconds = time.perf_counter() - started_clock
        with mp.get_context('spawn').Pool(args.workers, initializer=barcode.initialize_worker) as pool:
            for lot in lots:
                active = LotRun(lot.name, datetime.now().astimezone(), time.perf_counter())
                states.append(active)
                logger.info('Lot %s | START | %s', lot.name, active.started_at.isoformat(timespec='seconds'))
                print(f'Lot {lot.name} | START | {active.started_at.isoformat(timespec="seconds")}', flush=True)
                try:
                    active.work = Path(tempfile.mkdtemp(prefix=f'{lot.name}_', dir=incomplete))
                    records, page_rows, lot_matches, entries = process_lot(
                        lot, active, pool, database, targets, logger, None if args.legacy_layout else root)
                    logger.info('Lot %s | PUBLISH START | pages=%d', lot.name, active.processed_pages)
                    if args.legacy_layout:
                        save_lot_reports(active.work, records, page_rows)
                        publish_lot(active.work, output, lot.name)
                    else:
                        # The lot's report rows are kept only if all its PDFs are published.
                        with store.replace_lot(lot.name, run_id, entries):
                            publish_lot(active.work, output, lot.name,
                                        vis_output=root/'Output', ocr_output=root/'OCR')
                    active.published = True
                    matches.update(lot_matches)
                    if args.legacy_layout:
                        append_lot_csv(output / 'results.csv', records)
                        barcode.highlight_workbook(target_path, output / f'{target_path.stem}_found{target_path.suffix}', matches)
                    active.status = 'SAVED_WITH_ISSUES' if active.issues else 'SAVED'
                except Exception as exc:
                    active.status = 'FAILED'
                    message = f'Lot {lot.name}: {exc}'
                    active.issues.append(message)
                    logger.exception('%s | previously saved lots preserved', message)
                    print(message, file=sys.stderr, flush=True)
                    if active.work and active.work.exists():
                        logger.error('Lot %s | Unfinished files retained in %s', lot.name, active.work)
                issues.extend(active.issues)
                finish_lot(active, output, logger, lot_report=args.legacy_layout)
                if store is not None and active.published:
                    record_lot(root, store, run_id, active, logger)
                active = None
                update_report()
        if not any(state.published for state in states):
            raise barcode.ConfigurationError('No lots published: ' + '; '.join(issues))
        if args.legacy_layout:
            logger.info('Batch | Exporting combined pages.xlsx from saved lot reports')
            temporary = output / '.pages.xlsx.tmp'
            write_page_workbook(temporary, saved_page_rows(output, states))
            os.replace(temporary, output / 'pages.xlsx')
        finished_at = datetime.now().astimezone()
        elapsed = time.perf_counter() - started_clock
        duration = barcode.format_duration(elapsed)
        update_report('COMPLETED_WITH_ISSUES' if issues else 'COMPLETED', finished_at, elapsed)
        if all(state.published and not state.scan_errors for state in states):
            performance.record(output, 'barcode', workers=args.workers,
                               pages=sum(state.processed_pages for state in states), seconds=elapsed,
                               setup_seconds=setup_seconds, lots=len(states),
                               input_bytes=sum(state.input_bytes for state in states),
                               unresolved=sum(state.ocr for state in states), args=args)
        logger.info('Batch finished: %s | total processing time=%s (%.3f seconds) | lots saved=%d/%d',
                    finished_at.isoformat(timespec='seconds'), duration, elapsed,
                    sum(state.published for state in states), len(lots))
        print(f'Batch FINISHED | duration={duration} ({elapsed:.3f} seconds) | '
              f'lots saved={sum(state.published for state in states)}/{len(lots)} | '
              f'{output if args.legacy_layout else reporting.reports_dir(root)}', flush=True)
        return 1 if issues else 0
    except BaseException as exc:
        interrupted = isinstance(exc, KeyboardInterrupt)
        status = 'INTERRUPTED' if interrupted else 'FAILED'
        if active is not None and active.finished_at is None:
            active.status = status
            active.issues.append(f'Lot {active.name}: {status}; unfinished files retained in {active.work}')
            issues.extend(active.issues)
            finish_lot(active, output, logger, lot_report=args.legacy_layout)
        issues.append(f'Batch {status}: {exc or type(exc).__name__}')
        finished_at = datetime.now().astimezone()
        elapsed = time.perf_counter() - started_clock
        logger.error('Batch %s | duration=%s (%.3f seconds) | completed lots remain in %s',
                     status, barcode.format_duration(elapsed), elapsed, output)
        try:
            update_report(status, finished_at, elapsed)
        except Exception:
            logger.exception('Could not update final summary; saved PDFs and lot reports preserved')
        raise
    finally:
        if store is not None:
            # After Ctrl+C, stop now: the next stage or launch rewrites the stale reports.
            if not interrupted:
                print('Updating the cumulative reports…', flush=True)
                reporting.refresh_quietly(root, store, logger)
            store.close()
        split.close_logger(logger)
        if incomplete.exists() and not any(incomplete.iterdir()):
            incomplete.rmdir()


def record_lot(root: Path, store: reporting.Store, run_id: str, state: LotRun, logger) -> None:
    """Timing and remarks of a published lot, then its report folder."""
    try:
        store.record_lot_run(state.name, 'barcode', run_id, started=state.started_at,
                             finished=state.finished_at, seconds=state.elapsed,
                             status=state.status, notes=state.issues)
    except sqlite3.Error as exc:
        logger.error('Lot %s | timing not recorded in the reports: %s', state.name, exc)
    reporting.refresh_quietly(root, store, logger, complete=False)


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('-d', '--directory', type=Path, required=True)
    parser.add_argument('-n', '--workers', type=int, default=1)
    parser.add_argument('--database', type=Path, help='Override the BDD .xlsx discovered in the parent folder')
    parser.add_argument('--excel', type=Path, help='Search list directly inside the batch root')
    parser.add_argument('--legacy-layout', action='store_true', help='Use the old output folder layout')
    args = parser.parse_args(argv)
    if args.workers <= 0:
        parser.error('Workers must be greater than zero')
    return args


def main(argv=None):
    try:
        return run(parse_args(argv))
    except (barcode.ConfigurationError, split.ConfigurationError) as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print('Interrupted.', file=sys.stderr)
        return 130
    except Exception as exc:
        print(f'Fatal error: {type(exc).__name__}: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    mp.freeze_support()
    raise SystemExit(main())
