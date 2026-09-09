from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import ModuleType
from unittest.mock import patch

from scripts import script


class ClassificationTests(unittest.TestCase):
    def test_scanner_is_panorama_linear_code39_and_code128_only(self) -> None:
        captured: dict[str, object] = {}
        fake_codarascan = ModuleType("codarascan")

        class FakeScanner:
            def __init__(self, **kwargs: object) -> None:
                captured.update(kwargs)

        fake_codarascan.Scanner = FakeScanner
        with patch.dict("sys.modules", {"codarascan": fake_codarascan}):
            script.create_panorama_scanner()

        self.assertEqual(
            captured,
            {
                "mode": "panorama",
                "symbols": "linear",
                "formats": ("code-39", "code-128"),
                "decode": True,
            },
        )

    def test_exact_identifiers_only(self) -> None:
        self.assertEqual(script.classify_decoded_value(" 90432\n"), ("badge", "90432"))
        self.assertEqual(
            script.classify_decoded_value("vf1abc defT5702376"),
            None,
        )
        self.assertEqual(
            script.classify_decoded_value("vf1abcdeft5702376"),
            ("vin", "VF1ABCDEFT5702376"),
        )
        self.assertIsNone(script.classify_decoded_value("VF1ABCDEOT5702376"))
        self.assertIsNone(script.classify_decoded_value("prefix90432"))

    def test_badges_from_one_through_five_digits_are_supported(self) -> None:
        for badge in ("1", "10", "100", "1000", "90432"):
            with self.subTest(badge=badge):
                self.assertEqual(
                    script.classify_decoded_value(badge), ("badge", badge)
                )
                self.assertEqual(script.normalize_badge(badge), badge)

        self.assertIsNone(script.classify_decoded_value(""))
        self.assertIsNone(script.classify_decoded_value("123456"))
        self.assertIsNone(script.classify_decoded_value("-10"))

    def test_excel_badge_normalization_preserves_legacy_zero_padding(self) -> None:
        self.assertEqual(script.normalize_badge(10), "10")
        self.assertEqual(script.normalize_badge(1000.0), "1000")
        self.assertEqual(script.normalize_badge(10, "00000"), "00010")
        self.assertEqual(script.normalize_badge("0010"), "0010")

    def test_numeric_lot_sort(self) -> None:
        names = ["11", "2", "A", "1", "10"]
        self.assertEqual(
            sorted(names, key=script.lot_sort_key), ["1", "2", "10", "11", "A"]
        )


class PairingTests(unittest.TestCase):
    def setUp(self) -> None:
        ref1 = script.RowRef("Sheet1", 2)
        ref2 = script.RowRef("Sheet1", 3)
        self.targets = script.TargetIndex(
            workbook_path=Path("client.xlsx"),
            badge_rows={"90432": {ref1}, "91028": {ref2}},
            vis_rows={"T5702376": {ref1}, "T5704864": {ref2}},
        )

    def test_singletons_are_combined(self) -> None:
        page = script.PageResult(
            "1", 4, 10, badges=("90432",), vins=("VF1ABCDEFT5702376",)
        )
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].found, "FOUND")
        self.assertEqual(records[0].vis, "T5702376")

    def test_unexpected_plural_identifiers_still_use_one_page_row(self) -> None:
        page = script.PageResult(
            "1",
            4,
            10,
            badges=("90432", "91028"),
            vins=("VF1ABCDEFT5702376", "VF1ABCDEFT5704864"),
        )
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].badge, "90432;91028")
        self.assertEqual(
            records[0].vin, "VF1ABCDEFT5702376;VF1ABCDEFT5704864"
        )
        self.assertEqual(records[0].vis, "T5702376;T5704864")
        self.assertEqual(records[0].found, "FOUND")

    def test_single_badge_and_vin_stay_on_same_page_row_even_without_source_pair(
        self,
    ) -> None:
        page = script.PageResult(
            "1", 4, 10, badges=("90432",), vins=("VF1ABCDEFT5704864",)
        )
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].badge, "90432")
        self.assertEqual(records[0].vin, "VF1ABCDEFT5704864")
        self.assertEqual(records[0].vis, "T5704864")
        self.assertEqual(records[0].found, "FOUND")

    def test_unknown_single_badge_and_vin_stay_on_same_page_row(self) -> None:
        page = script.PageResult(
            "1", 5, 10, badges=("91010",), vins=("BRYEKNFJXT5704852",)
        )
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(
            records,
            [
                script.CsvRecord(
                    "1",
                    5,
                    badge="91010",
                    vin="BRYEKNFJXT5704852",
                    vis="T5704852",
                    found="",
                )
            ],
        )

    def test_page_without_identifiers_gets_one_empty_csv_row(self) -> None:
        page = script.PageResult("1", 6, 10)
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(records, [script.CsvRecord("1", 6)])

    def test_failed_page_can_be_represented_by_one_empty_csv_row(self) -> None:
        page = script.PageResult("1", 7, 10, error="render failed")
        records = script.csv_records_for_page(page, self.targets)
        self.assertEqual(records, [script.CsvRecord("1", 7)])


class RenderingTests(unittest.TestCase):
    def test_render_is_300_dpi_then_cropped_to_exact_upper_half(self) -> None:
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow is unavailable")

        calls: dict[str, object] = {}

        class FakeBitmap:
            closed = False

            def to_pil(self):
                return Image.new("RGB", (101, 201), "white")

            def close(self) -> None:
                self.closed = True

        class FakePage:
            closed = False

            def render(self, *, scale: float, rotation: int):
                calls["scale"] = scale
                calls["rotation"] = rotation
                calls["bitmap"] = FakeBitmap()
                return calls["bitmap"]

            def close(self) -> None:
                self.closed = True

        class FakeDocument:
            def __getitem__(self, page_index: int):
                calls["page_index"] = page_index
                calls["page"] = FakePage()
                return calls["page"]

        roi = script.render_upper_half(FakeDocument(), 3)
        try:
            self.assertEqual(roi.size, (101, 100))
            self.assertAlmostEqual(calls["scale"], 300 / 72)
            self.assertEqual(calls["rotation"], 0)
            self.assertEqual(calls["page_index"], 3)
            self.assertTrue(calls["bitmap"].closed)
            self.assertTrue(calls["page"].closed)
        finally:
            roi.close()

    def test_page_failure_is_returned_instead_of_raised(self) -> None:
        class BrokenScanner:
            def scan_image(self, image, *, diagnostics: bool):
                raise RuntimeError("deliberate scan failure")

        class FakeRoi:
            closed = False

            def close(self) -> None:
                self.closed = True

        old_scanner = script._WORKER_SCANNER
        roi = FakeRoi()
        script._WORKER_SCANNER = BrokenScanner()
        try:
            with (
                patch.object(script, "_worker_document", return_value=object()),
                patch.object(script, "render_upper_half", return_value=roi),
            ):
                result = script.process_page(script.PageTask("7", "bad.pdf", 2, 12))
        finally:
            script._WORKER_SCANNER = old_scanner

        self.assertEqual(result.page_number, 3)
        self.assertIn("deliberate scan failure", result.error or "")
        self.assertIn("RuntimeError", result.traceback_text or "")
        self.assertTrue(roi.closed)


class WorkbookTests(unittest.TestCase):
    def test_reference_workbook_structure(self) -> None:
        reference = Path(
            "/Users/fellahiabderraouf/Downloads/Dirty List VIN Delivered.xlsx"
        )
        if not reference.exists():
            self.skipTest("reference workbook is unavailable")
        targets = script.load_target_index(reference)
        self.assertEqual(len(targets.rows), 229)
        self.assertIn("90432", targets.badge_rows)
        self.assertIn("T5702376", targets.vis_rows)

    def test_highlight_preserves_values_and_marks_only_found_row(self) -> None:
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("openpyxl is unavailable")
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "client.xlsx"
            output = Path(directory) / "client_found.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["BDG", "VIS", "Note"])
            sheet.append(["90432", "T5702376", "keep me"])
            sheet.append(["91028", "T5704864", "unchanged"])
            workbook.save(source)
            script.highlight_workbook(source, output, {script.RowRef(sheet.title, 2)})
            result = load_workbook(output)
            self.assertEqual(result.active["C2"].value, "keep me")
            self.assertEqual(result.active["A2"].fill.fgColor.rgb, "00C6EFCE")
            self.assertNotEqual(result.active["A3"].fill.fgColor.rgb, "00C6EFCE")


if __name__ == "__main__":
    unittest.main()
