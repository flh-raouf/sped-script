from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from test_split_pages import make_pdf

import main
from scripts import classify_pages as classify
from scripts import ocr_pages as ocr

VIN = 'BRYEKNFJ1T5704898'
VIS = VIN[-8:]


def line(text, box=(10, 10, 200, 30), score=.99):
    return ocr.TextLine(text, score, box)


def database():
    record = classify.VehicleRecord('91034', VIN, VIS, 'SEQEMON0118300161', '6G3D1411', 2026)
    other = classify.VehicleRecord('91069', 'BRYEKNFJ5S5780509', 'S5780509', 'SEQEMON0118300182', '', 2026)
    return classify.DatabaseIndex({'91034': frozenset({VIS}), '91069': frozenset({'S5780509'})},
                                  {VIS: [record], 'S5780509': [other]})


class CacheTests(unittest.TestCase):
    def test_cache_survives_restart_and_distinguishes_content_and_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'cache.sqlite3'
            cache = ocr.OcrCache(path, {'dpi': 200, 'model': 'tiny'})
            value = [line('VIN: ' + VIN)]
            cache.put('digest1', value)
            cache.put('empty', [])
            cache.close()
            cache = ocr.OcrCache(path, {'model': 'tiny', 'dpi': 200})
            self.assertEqual(cache.get('digest1'), value)
            self.assertEqual(cache.get('empty'), [])
            self.assertIsNone(cache.get('digest2'))
            cache.close()
            cache = ocr.OcrCache(path, {'dpi': 300, 'model': 'tiny'})
            self.assertIsNone(cache.get('digest1'))
            cache.close()

    def test_malformed_cache_is_a_miss_not_an_empty_success(self):
        with tempfile.TemporaryDirectory() as directory:
            cache = ocr.OcrCache(Path(directory) / 'cache.sqlite3', {})
            for payload in ('{}', 'null', '[{"text":"VIN","score":0.9,"box":[1,2]}]',
                            '[{"text":"VIN","score":2,"box":[1,2,3,4]}]'):
                cache.connection.execute('INSERT OR REPLACE INTO ocr VALUES (?,?)', (cache.key('x'), payload))
                self.assertIsNone(cache.get('x'))
            cache.close()


class ResolutionTests(unittest.TestCase):
    def setUp(self):
        self.lookup = ocr.Lookup(database())

    def test_vin_and_labelled_badge_identify(self):
        for text in (VIN, 'BDG: 91034', 'Badge: 91034'):
            self.assertEqual(ocr.resolve_text([line(text)], self.lookup).vis, VIS)
        for text in ('91034', 'SEQ: 999999999'):
            self.assertIsNone(ocr.resolve_text([line(text)], self.lookup).vis)

    def test_direct_vis_known_to_bdd_identifies_and_enriches_the_page(self):
        for text in (VIS, 'VIS: ' + VIS, 'VIS #:' + VIS):
            decision = ocr.resolve_text([line(text)], self.lookup)
            self.assertEqual(decision.vis, VIS)
            self.assertEqual(ocr.enriched_row('1_3', decision, self.lookup.database),
                             ('1_3', '91034', VIN, VIS, 'SEQEMON0118300161', '118300161', '6G3D1411'))
        self.assertEqual(ocr.resolve_text([line('VIS'), line(VIS)], self.lookup).vis, VIS)

    def test_direct_vis_does_not_accept_unknown_low_score_or_embedded_values(self):
        for text in ('VIS: T5999999', 'KT5704898JOS0590E01210G404C', 'XT5704898', VIS+'X'):
            self.assertIsNone(ocr.resolve_text([line(text)], self.lookup).vis)
        self.assertIsNone(ocr.resolve_text([line('VIS: '+VIS, score=.5)], self.lookup).vis)

    def test_direct_vis_conflicts_with_other_reliable_identifiers(self):
        for other in ('VIS: S5780509', 'BDG: 91069', 'SEQ: 126183182', 'BRYEKNFJ5S5780509'):
            self.assertEqual(ocr.resolve_text([line('VIS: '+VIS), line(other)], self.lookup).reason,
                             'IDENTIFIER_CONFLICT')

    def test_separate_labels_use_positions_and_reject_operator_badges(self):
        label = line('BDG:', (10, 10, 60, 30))
        value = line('91034', (65, 10, 150, 30))
        self.assertEqual(ocr.resolve_text([label, value], self.lookup).vis, VIS)
        self.assertIsNone(ocr.resolve_text([label, line('91034', (10, 500, 100, 520))], self.lookup).vis)
        self.assertIsNone(ocr.resolve_text([line('Badge opérateur: 91034')], self.lookup).vis)

    def test_sequence_spaces_last_nine_and_document_alias(self):
        for text in ('Séquence: 1 26 183 161', 'SEQ: 118300161'):
            self.assertEqual(ocr.resolve_text([line(text)], self.lookup).vis, VIS)
        self.assertEqual(ocr.resolve_text([line('91034'), line('126183161')], self.lookup).vis, VIS)
        self.assertIsNone(ocr.resolve_text([line('126183161')], self.lookup).vis)
        # Other BDD sequence formats still support the requested last-nine lookup.
        record = classify.VehicleRecord('00123', VIN, VIS, '50FV750960305')
        lookup = ocr.Lookup(classify.DatabaseIndex({'00123': frozenset({VIS})}, {VIS: [record]}))
        self.assertEqual(ocr.resolve_text([line('SEQ: 750960305')], lookup).vis, VIS)

    def test_conflicting_identifiers_low_scores_and_ambiguous_sequences_go_to_review(self):
        self.assertEqual(ocr.resolve_text([line(VIN), line('BDG: 91069')], self.lookup).reason,
                         'IDENTIFIER_CONFLICT')
        self.assertIsNone(ocr.resolve_text([line(VIN, score=.5)], self.lookup).vis)
        self.lookup.sequences['126183161'].add('S5780509')
        self.assertEqual(ocr.resolve_text([line('SEQ: 126183161')], self.lookup).reason,
                         'AMBIGUOUS_IDENTIFIER')

    def test_sequence_enrichment_selects_the_matching_record_for_a_shared_vis(self):
        db = database()
        db.records_by_vis[VIS].append(classify.VehicleRecord('11111', VIN, VIS, 'SEQEMON0118300999', 'OTHER', 2026))
        lookup = ocr.Lookup(db)
        decision = ocr.resolve_text([line('SEQ: 126183161')], lookup)
        self.assertEqual(ocr.enriched_row('479_3', decision, db),
                         ('479_3', '91034', VIN, VIS, 'SEQEMON0118300161', '118300161', '6G3D1411'))


class TransferTests(unittest.TestCase):
    def test_moves_and_keeps_existing_vis_pages(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, destination = root / '1_2.pdf', root / VIS / '1_2.pdf'
            source.write_bytes(b'page')
            ocr.transfer_page(source, destination, ocr.file_digest(source), False)
            self.assertFalse(source.exists())
            self.assertEqual(destination.read_bytes(), b'page')

    def test_collision_and_concurrent_creation_preserve_both_documents(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, destination = root / 'input.pdf', root / 'output.pdf'
            source.write_bytes(b'source')
            destination.write_bytes(b'existing')
            with self.assertRaises(FileExistsError):
                ocr.transfer_page(source, destination, ocr.file_digest(source), False)
            destination.unlink()
            original = Path.open
            def concurrent_open(path, mode='r', *args, **kwargs):
                if path == destination and mode == 'xb':
                    destination.write_bytes(b'concurrent')
                return original(path, mode, *args, **kwargs)
            with patch.object(Path, 'open', concurrent_open), self.assertRaises(FileExistsError):
                ocr.transfer_page(source, destination, ocr.file_digest(source), False)
            self.assertEqual(source.read_bytes(), b'source')
            self.assertEqual(destination.read_bytes(), b'concurrent')


class BatchTests(unittest.TestCase):
    def create_batch(self, root):
        from openpyxl import Workbook
        source = root / 'unresolved'
        source.mkdir()
        make_pdf(source / '1_1.pdf', ['VIN page'])
        make_pdf(source / '1_2.pdf', ['Unidentified page'])
        bdd = root / 'BDD_24_26.xlsx'
        workbook = Workbook()
        workbook.active.append(['Année', 'BDG', 'VIS', 'SEM', 'ID2', 'NOF'])
        workbook.active.append([2026, '91034', VIS, 'SEQEMON0118300161', VIN, '6G3D1411'])
        workbook.save(bdd)
        workbook.close()
        return source, bdd

    def test_progressive_outputs_cache_hits_and_changed_bdd_reenrichment(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, bdd = self.create_batch(root)
            args = ocr.parse_args(['--legacy-layout', '-d', str(source), '--copy'])
            def read(path):
                if path.stem == '1_2':
                    self.assertTrue((root / VIS / '1_1.pdf').is_file())
                    self.assertTrue(next((root / 'ocr_runs').glob('*/results.csv')).stat().st_size > 100)
                    return []
                return [line('BDG: 91034')]
            with patch.object(ocr, 'PaddleEngine') as engine:
                engine.return_value.read.side_effect = read
                self.assertEqual(ocr.run(args), 0)
                self.assertEqual(engine.return_value.read.call_count, 2)
            self.assertTrue((root / 'Review' / '1_2.pdf').is_file())
            workbook = load_workbook(bdd)
            workbook.active['F2'] = 'NEWNOF'
            workbook.save(bdd)
            workbook.close()
            with patch.object(ocr, 'PaddleEngine', side_effect=AssertionError('OCR must not run')):
                self.assertEqual(ocr.run(args), 0)
            reports = max((root / 'ocr_runs').iterdir())
            with (reports / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([row['Cache'] for row in rows], ['HIT', 'HIT'])
            self.assertEqual(rows[0]['NOF'], 'NEWNOF')
            self.assertEqual(rows[1]['VIS'], '')
            self.assertIn('Total duration:', (reports / 'report.txt').read_text())
            self.assertTrue((reports / 'pages.xlsx').is_file())

    def test_failed_ocr_is_not_cached_and_next_page_is_processed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, _ = self.create_batch(root)
            args = ocr.parse_args(['--legacy-layout', '-d', str(source), '--copy'])
            with patch.object(ocr, 'PaddleEngine') as engine:
                engine.return_value.read.side_effect = [RuntimeError('scan failed'), []]
                self.assertEqual(ocr.run(args), 1)
            cache = ocr.OcrCache(root / '.ocr_cache.sqlite3', ocr.ocr_settings(200, 'cpu'))
            self.assertIsNone(cache.get(ocr.file_digest(source / '1_1.pdf')))
            self.assertEqual(cache.get(ocr.file_digest(source / '1_2.pdf')), [])
            cache.close()

    def test_menu_launches_ocr_with_explicit_destination(self):
        command = main.build_command('ocr_pages.py', Path('/unresolved'), output=Path('/vis-root'))
        self.assertEqual(command[-2:], ['--output', '/vis-root'])
        self.assertNotIn('-n', command)


if __name__ == '__main__':
    unittest.main()
