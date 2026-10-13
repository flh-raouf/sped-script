from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from test_ocr_pages import VIN, VIS, line
from test_review_pages import parts_table, stage_rows
from test_split_pages import make_pdf

from scripts import classify_pages as classify
from scripts import ocr_pages as ocr
from scripts import reporting
from scripts import review_pages as review
from scripts import performance


def sheet_rows(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        header, *rows = list(workbook.active.values)
        return [dict(zip(header, row)) for row in rows]
    finally:
        workbook.close()


class ParentWorkflowTests(unittest.TestCase):
    def test_full_sequence_uses_same_parent_shared_cache_and_preserves_originals(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'1').mkdir()
            original = root/'1/lot.pdf'
            make_pdf(original, ['barcode front', 'OCR front', 'back', 'table', 'unknown'])
            original_bytes = original.read_bytes()
            workbook = Workbook()
            workbook.active.append(['Année', 'BDG', 'VIS', 'SEM', 'ID2', 'NOF'])
            workbook.active.append([2026, '91034', VIS, 'SEQEMON0118300161', VIN, '6G3D1411'])
            workbook.save(root/'BDD.xlsx')
            workbook.close()
            workbook = Workbook()
            workbook.active.append(['BDG', 'VIS'])
            workbook.active.append(['91034', VIS])
            workbook.save(root/'search.xlsx')
            workbook.close()

            class Pool:
                def __enter__(self):
                    return self

                def __exit__(self, *args):
                    pass

                def imap_unordered(self, function, tasks, chunksize):
                    return (function(task) for task in tasks)

            def scan(task):
                return ocr.barcode.PageResult(task.lot, task.page_index+1, task.pages_in_lot,
                                              badges=('91034',) if Path(task.pdf_path).stem == '1_1' else ())

            with patch.object(ocr.barcode, 'validate_runtime'), \
                 patch.object(ocr.barcode, 'process_page', side_effect=scan), \
                 patch.object(ocr.barcode, '_close_worker_document'), \
                 patch.object(classify.mp, 'get_context', return_value=SimpleNamespace(Pool=lambda *a, **k: Pool())):
                self.assertEqual(classify.run(classify.parse_args(['-d', str(root)])), 0)
            self.assertTrue((root/'Output'/VIS/'1_1.pdf').exists())
            self.assertEqual(len(list((root/'OCR').glob('*.pdf'))), 4)
            self.assertFalse((root/'Output/OCR').exists())
            reports = root/'Rapports'
            pages = {row['lot_page']: row for row in sheet_rows(reports/'Pages/pages_lots_0001-0200.xlsx')}
            self.assertEqual(list(pages), ['1_1', '1_2', '1_3', '1_4', '1_5'])
            self.assertEqual((pages['1_1']['État'], pages['1_1']['Classée par'], pages['1_1']['VIS'],
                              pages['1_1']['Liste de recherche'], pages['1_1']['PDF']),
                             ('Classée', 'Code-barres', VIS, 'Trouvée', f'Output/{VIS}/1_1.pdf'))
            self.assertEqual(pages['1_2']['État'], "En attente d'OCR")
            self.assertEqual([row['lot_page'] for row in sheet_rows(reports/'Lots/1/pages.xlsx')], list(pages))
            self.assertTrue((reports/'Lots/1/code-barres.csv').is_file())
            self.assertIn('Début :', (reports/'Lots/1/rapport.txt').read_text(encoding='utf-8'))
            [lot] = sheet_rows(reports/'suivi_lots.xlsx')
            self.assertEqual((lot['Lot'], lot['Pages'], lot['Classées'], lot["En attente d'OCR"], lot['État']),
                             ('1', 5, 1, 4, 'En cours'))
            from openpyxl import load_workbook
            found = load_workbook(reports/'search_found.xlsx')
            self.assertEqual(found.active['A2'].fill.fgColor.rgb, '00C6EFCE')
            found.close()
            self.assertIn('Batch status: COMPLETED',
                          next((reports/'Executions').glob('* - code-barres.txt')).read_text())

            def read(path):
                return {'1_2': [line('BDG: 91034')], '1_3': [],
                        '1_4': parts_table(), '1_5': [line('CHECK-LIST')]}[path.stem]
            with patch.object(ocr, 'PaddleEngine') as engine:
                engine.return_value.read.side_effect = read
                self.assertEqual(ocr.run(ocr.parse_args(['-d', str(root)])), 0)
                self.assertEqual(engine.return_value.read.call_count, 4)
            self.assertEqual(len(list((root/'OCR').glob('*.pdf'))), 0)
            self.assertEqual(len(list((root/'Pending').glob('*.pdf'))), 3)
            self.assertEqual(reporting.cache_path(root), reports/'.interne/cache_ocr.sqlite3')
            self.assertTrue(reporting.cache_path(root).exists())
            self.assertFalse((root/'Review').exists())
            self.assertTrue((reports/'Lots/1/ocr.csv').is_file())

            with patch.object(ocr, 'PaddleEngine', side_effect=AssertionError('Use shared OCR cache')), \
                 patch.object(review, 'ink_percent', side_effect=lambda path, **_: .1 if path.stem == '1_3' else 5):
                self.assertEqual(review.run(review.parse_args(['-d', str(root), '--cache-only'])), 0)
            self.assertEqual(len(list((root/'Pending').glob('*.pdf'))), 0)
            self.assertEqual(len(list((root/'Output'/VIS).glob('*.pdf'))), 4)
            self.assertTrue((root/'Review/1_5.pdf').exists())
            self.assertTrue((reports/'.interne/review_assignments.jsonl').exists())
            for stage, pages, unresolved in [('barcode', 5, 4), ('ocr', 4, 3), ('review', 3, 1)]:
                settings = performance.options(stage)
                if stage == 'review':
                    settings['cache_only'] = True
                history = performance.history(root, stage, settings)
                self.assertEqual(len(history), 1, stage)
                self.assertEqual((history[0].pages, history[0].unresolved), (pages, unresolved))
            rows = {row['lot_page']: row for row in stage_rows(root, 'review')}
            self.assertEqual(rows['1_3']['BDG'], '91034')
            self.assertEqual(rows['1_3']['Predecessor'], '1_2')
            self.assertEqual(Path(rows['1_5']['PDF']), (root/'Review/1_5.pdf').resolve())
            # Each stage completed the same cumulative rows rather than adding new files.
            pages = {row['lot_page']: row for row in sheet_rows(reports/'Pages/pages_lots_0001-0200.xlsx')}
            self.assertEqual({name: (row['État'], row['Classée par']) for name, row in pages.items()}, {
                '1_1': ('Classée', 'Code-barres'), '1_2': ('Classée', 'OCR'),
                '1_3': ('Classée', 'Classification finale'), '1_4': ('Classée', 'Classification finale'),
                '1_5': ('À revoir', None)})
            self.assertEqual((pages['1_3']['BDG'], pages['1_3']['Classification'], pages['1_5']['PDF']),
                             ('91034', 'PREVIOUS_PAGE', 'Review/1_5.pdf'))
            self.assertEqual((pages['1_2']['Code-barres'], pages['1_2']['OCR']), ('NO_BARCODE', 'BDD_MATCH'))
            report = (reports/'Lots/1/rapport.txt').read_text(encoding='utf-8')
            self.assertIn('Classées : 4 (code-barres : 1 · OCR : 1 · classification finale : 2)', report)
            self.assertIn('5 pages : 1 classées, 4 envoyées en OCR', report)
            self.assertIn('4 pages : 1 classées, 3 envoyées en classification finale, 0 en erreur', report)
            self.assertIn('3 pages : 2 classées, 1 à revoir, 0 en erreur', report)
            with (reports/'Lots/1/classification.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual([row['lot_page'] for row in csv.DictReader(handle)], ['1_3', '1_4', '1_5'])
            [lot] = sheet_rows(reports/'suivi_lots.xlsx')
            self.assertEqual((lot['Classées'], lot['À revoir'], lot['Pages trouvées dans la liste'], lot['État']),
                             (4, 1, 4, 'Terminé'))
            summary = (reports/'resume.txt').read_text(encoding='utf-8')
            self.assertIn('Classées : 4 (80,0 %)', summary)
            self.assertIn('À revoir (dossier Review) : 1', summary)
            self.assertIn('« search.xlsx » : 1 lignes trouvées sur 1', summary)
            self.assertEqual(reporting.pending_updates(root), 0)
            self.assertEqual(original.read_bytes(), original_bytes)
            # Empty stages are safe no-ops and don't initialize OCR again.
            self.assertEqual(ocr.run(ocr.parse_args(['-d', str(root)])), 0)
            self.assertEqual(review.run(review.parse_args(['-d', str(root)])), 0)
            # Re-running barcode classification cannot clear or overwrite results.
            saved = {path.name: path.read_bytes() for path in (root/'Output'/VIS).glob('*.pdf')}
            with patch.object(ocr.barcode, 'validate_runtime'), \
                 patch.object(ocr.barcode, 'process_page', side_effect=scan), \
                 patch.object(ocr.barcode, '_close_worker_document'), \
                 patch.object(classify.mp, 'get_context', return_value=SimpleNamespace(Pool=lambda *a, **k: Pool())), \
                 self.assertRaisesRegex(ocr.barcode.ConfigurationError, 'No lots published'):
                classify.run(classify.parse_args(['-d', str(root)]))
            self.assertEqual(saved, {path.name: path.read_bytes() for path in (root/'Output'/VIS).glob('*.pdf')})
            self.assertFalse(list((root/'OCR').glob('*.pdf')))
            # The refused rescan left the lot's recorded results untouched.
            self.assertEqual(sheet_rows(reports/'Pages/pages_lots_0001-0200.xlsx'), list(pages.values()))

    def test_missing_stage_folder_explains_previous_step_without_creating_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(ValueError, 'Run barcode classification first'):
                ocr.run(ocr.parse_args(['-d', str(root)]))
            with self.assertRaisesRegex(ValueError, 'Run OCR classification first'):
                review.run(review.parse_args(['-d', str(root)]))
            self.assertFalse((root/'Output').exists())

    def test_cache_miss_in_last_stage_moves_to_final_review(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'Pending').mkdir()
            make_pdf(root/'Pending/1_1.pdf', ['unknown'])
            from test_ocr_pages import database
            with patch.object(ocr, 'discover_bdd', return_value=root/'BDD.xlsx'), \
                 patch.object(classify, 'load_database', return_value=database()), \
                 patch.object(ocr, 'PaddleEngine', side_effect=AssertionError('No OCR')):
                self.assertEqual(review.run(review.parse_args(['-d', str(root), '--cache-only'])), 0)
            self.assertTrue((root/'Review/1_1.pdf').exists())
            self.assertFalse((root/'Pending/1_1.pdf').exists())

    def test_empty_stages_still_create_the_next_folder_without_requiring_bdd_or_ocr(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'OCR').mkdir()
            self.assertEqual(ocr.run(ocr.parse_args(['-d', str(root)])), 0)
            self.assertTrue((root/'Pending').is_dir())
            self.assertEqual(review.run(review.parse_args(['-d', str(root)])), 0)
            self.assertTrue((root/'Review').is_dir())


if __name__ == '__main__':
    unittest.main()
