from __future__ import annotations

import csv
import json
import os
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from scripts import reporting
from scripts.reporting import CLASSIFIED, TO_REVIEW, WAITING_OCR, WAITING_REVIEW, PageEntry, Store

VIS = 'T5704898'


def sheet_rows(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        header, *rows = list(workbook.active.values)
        return [dict(zip(header, row)) for row in rows]
    finally:
        workbook.close()


def scanned(lot, page, vis=None, badges=()):
    row = (f'{lot}_{page}', badges[0] if badges else '', '', vis or '', '', '', '')
    return PageEntry(lot, page, CLASSIFIED if vis else WAITING_OCR, row, 'BDD_BADGE' if vis else 'NO_BARCODE',
                     {'Lot': lot, 'Page': str(page)}, pdf=f'Output/{vis}/{lot}_{page}.pdf' if vis else f'OCR/{lot}_{page}.pdf',
                     badges=tuple(badges))


def write_list(path, rows):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['BDG', 'VIS'])
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


class ReportingTestCase(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.store = Store(self.root)
        self.addCleanup(self.store.close)
        self.started = datetime(2026, 10, 13, 10, 0).astimezone()

    def scan(self, lot, pages, run='20261013_100000_000000', vis_pages=()):
        entries = [scanned(lot, page, VIS if page in vis_pages else None) for page in range(1, pages + 1)]
        with self.store.replace_lot(lot, run, entries):
            pass
        self.store.record_lot_run(lot, 'barcode', run, started=self.started, finished=self.started,
                                  seconds=12, status='SAVED')

    def global_pages(self, name='pages_lots_0001-0200.xlsx'):
        return sheet_rows(self.root / 'Rapports' / 'Pages' / name)


class CumulativeReportTests(ReportingTestCase):
    def test_each_run_adds_its_lots_to_the_same_global_files_in_lot_order(self):
        self.scan('2', 2, vis_pages={1})
        self.assertEqual(reporting.refresh(self.root, self.store), ['Aucune liste de recherche connue : '
                                                                     "la liste coloriée n'est pas mise à jour."])
        self.scan('1', 3, run='20261014_100000_000000')
        reporting.refresh(self.root, self.store)
        self.assertEqual([row['lot_page'] for row in self.global_pages()], ['1_1', '1_2', '1_3', '2_1', '2_2'])
        self.assertEqual([row['Lot'] for row in sheet_rows(self.root / 'Rapports' / 'suivi_lots.xlsx')], ['1', '2'])
        self.assertTrue((self.root / 'Rapports' / 'Lots' / '1' / 'pages.xlsx').is_file())
        self.assertTrue((self.root / 'Rapports' / 'Lots' / '2' / 'rapport.txt').is_file())
        self.assertIn('Lots : 2 (0 terminés, 2 en cours)',
                      (self.root / 'Rapports' / 'resume.txt').read_text(encoding='utf-8'))

    def test_lots_are_split_into_parts_of_200_lots(self):
        for lot in ('1', '200', '201', '0'):
            self.scan(lot, 1)
        reporting.refresh(self.root, self.store)
        names = sorted(path.name for path in (self.root / 'Rapports' / 'Pages').iterdir())
        self.assertEqual(names, ['pages_lots_0001-0200.xlsx', 'pages_lots_0201-0400.xlsx'])
        self.assertEqual([row['lot_page'] for row in self.global_pages()], ['0_1', '1_1', '200_1'])
        self.assertEqual(reporting.part_name(49), 'pages_lots_09801-10000.xlsx')

    def test_only_the_parts_and_lots_touched_are_rewritten(self):
        self.scan('1', 1)
        self.scan('201', 1)
        reporting.refresh(self.root, self.store)
        self.scan('202', 1)
        with patch.object(reporting, 'write_part', wraps=reporting.write_part) as parts, \
                patch.object(reporting, 'write_lot', wraps=reporting.write_lot) as lots:
            reporting.refresh(self.root, self.store)
        self.assertEqual([call.args[2] for call in parts.call_args_list], [1])
        self.assertEqual([call.args[2] for call in lots.call_args_list], ['202'])

    def test_rescanning_a_lot_replaces_its_rows(self):
        self.scan('1', 3)
        self.store.record('ocr', 'run', PageEntry('1', 2, WAITING_REVIEW, ('1_2',) + ('',) * 6, 'NO_RELIABLE_IDENTIFIER'))
        self.scan('1', 2, run='20261015_100000_000000', vis_pages={1})
        reporting.refresh(self.root, self.store)
        rows = self.global_pages()
        self.assertEqual([(row['lot_page'], row['État'], row['OCR']) for row in rows],
                         [('1_1', 'Classée', None), ('1_2', "En attente d'OCR", None)])
        self.assertFalse((self.root / 'Rapports' / 'Lots' / '1' / 'ocr.csv').exists())

    def test_a_failed_publication_keeps_the_earlier_scan(self):
        self.scan('1', 2)
        with self.assertRaises(FileExistsError), self.store.replace_lot('1', 'again', [scanned('1', 1)]):
            raise FileExistsError('collision')
        self.assertEqual(self.store.lot_counts('1')[0][1], 2)

    def test_a_failed_page_keeps_what_earlier_stages_found_and_corrects_its_intended_move(self):
        self.scan('1', 1)
        intent = PageEntry('1', 1, CLASSIFIED, ('1_1', '91034', '', VIS, '', '', ''), 'BDD_MATCH',
                           {'Status': 'BDD_MATCH'}, pdf=f'Output/{VIS}/1_1.pdf', badges=('91034',), seconds=2)
        recorded = self.store.record('ocr', 'run', intent)
        failure = PageEntry('1', 1, WAITING_OCR, ('1_1',) + ('',) * 6, 'ERROR', {'Status': 'ERROR'},
                            pdf='OCR/1_1.pdf', error='FileExistsError: collision', seconds=2)
        self.store.record('ocr', 'run', failure, replaces=recorded)
        reporting.refresh(self.root, self.store)
        [row] = self.global_pages()
        self.assertEqual((row['État'], row['OCR'], row['Erreur'], row['PDF'], row['BDG']),
                         ("En attente d'OCR", 'ERROR', 'FileExistsError: collision', 'OCR/1_1.pdf', '91034'))
        with (self.root / 'Rapports' / 'Lots' / '1' / 'ocr.csv').open(encoding='utf-8-sig', newline='') as handle:
            self.assertEqual([row['Status'] for row in csv.DictReader(handle)], ['ERROR'])
        report = (self.root / 'Rapports' / 'Lots' / '1' / 'rapport.txt').read_text(encoding='utf-8')
        self.assertIn('1 pages : 0 classées, 0 envoyées en classification finale, 1 en erreur', report)

    def test_review_metadata_comes_from_classified_pages(self):
        self.scan('1', 2, vis_pages={1})
        self.assertEqual(self.store.metadata(), {'1_1': ('1_1', '', '', VIS, '', '', '')})


class SearchListTests(ReportingTestCase):
    def test_list_highlights_every_page_found_so_far_and_follows_a_replaced_list(self):
        from openpyxl import load_workbook
        first = self.root / 'dirty list.xlsx'
        write_list(first, [['91034', VIS], ['55555', 'T0000001']])
        self.store.set_search_list(first)
        entries = [scanned('1', 1, badges=('91034',)), scanned('1', 2, VIS)]
        with self.store.replace_lot('1', 'run', entries):
            pass
        reporting.refresh(self.root, self.store)
        workbook = load_workbook(self.root / 'Rapports' / 'dirty list_found.xlsx')
        fills = [workbook.active.cell(row, 1).fill.fgColor.rgb for row in (2, 3)]
        workbook.close()
        self.assertEqual(fills, ['00C6EFCE', '00000000'])
        self.assertEqual([row['Liste de recherche'] for row in self.global_pages()], ['Trouvée', 'Trouvée'])
        self.assertIn('1 lignes trouvées sur 2', (self.root / 'Rapports' / 'resume.txt').read_text(encoding='utf-8'))
        replacement = self.root / 'new list.xlsx'
        write_list(replacement, [['77777', 'T9999999']])
        self.store.mark_everything()
        reporting.refresh(self.root, self.store, search_list=replacement)
        self.assertEqual([row['Liste de recherche'] for row in self.global_pages()], [None, None])
        self.assertTrue((self.root / 'Rapports' / 'new list_found.xlsx').is_file())

    def test_damaged_list_is_reported_without_stopping_the_reports(self):
        broken = self.root / 'list.xlsx'
        broken.write_bytes(b'not a workbook')
        self.store.set_search_list(broken)
        self.scan('1', 1)
        problems = reporting.refresh(self.root, self.store)
        self.assertTrue(any('illisible' in problem for problem in problems))
        self.assertTrue((self.root / 'Rapports' / 'suivi_lots.xlsx').is_file())


class LockedFileTests(ReportingTestCase):
    def test_a_file_open_in_excel_is_retried_by_the_next_refresh(self):
        self.scan('1', 1)
        target = self.root / 'Rapports' / 'Pages' / 'pages_lots_0001-0200.xlsx'
        original = os.replace
        def locked(source, destination):
            if Path(destination) == target:
                raise PermissionError(13, 'Permission denied', str(destination))
            return original(source, destination)
        with patch.object(reporting.os, 'replace', side_effect=locked):
            problems = reporting.refresh(self.root, self.store)
        self.assertEqual(len([problem for problem in problems if 'Pages/pages_lots_0001-0200.xlsx' in problem]), 1)
        self.assertFalse(target.exists())
        self.assertFalse(list(target.parent.glob('.*')))   # no temporary file left behind
        self.assertIn('Non mis à jour', (self.root / 'Rapports' / 'resume.txt').read_text(encoding='utf-8'))
        self.assertEqual(reporting.pending_updates(self.root), 1)
        reporting.refresh(self.root, self.store)
        self.assertTrue(target.is_file())
        self.assertEqual(reporting.pending_updates(self.root), 0)

    def test_quiet_refresh_never_raises(self):
        with patch.object(reporting, 'refresh', side_effect=sqlite3.OperationalError('disk I/O error')):
            problems = reporting.refresh_quietly(self.root, self.store)
        self.assertIn('disk I/O error', problems[0])


class LayoutTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)

    def test_state_of_the_former_reports_folder_is_taken_over(self):
        legacy = self.root / 'Reports'
        legacy.mkdir()
        cache = sqlite3.connect(legacy / '.ocr_cache.sqlite3')
        cache.execute('PRAGMA journal_mode=WAL')
        cache.execute('CREATE TABLE ocr (key TEXT PRIMARY KEY, payload TEXT NOT NULL)')
        cache.execute("INSERT INTO ocr VALUES ('page', '[]')")
        cache.commit()
        (legacy / 'review_assignments.jsonl').write_text('{"name": "1_2"}\n')
        (legacy / 'pipeline_state.json').write_text(json.dumps({'version': 1, 'lots': {}}))
        self.assertEqual(reporting.cache_path(self.root), legacy / '.ocr_cache.sqlite3')
        cache.close()
        reporting.prepare(self.root)
        internal = self.root / 'Rapports' / '.interne'
        self.assertEqual(reporting.cache_path(self.root), internal / 'cache_ocr.sqlite3')
        moved = sqlite3.connect(internal / 'cache_ocr.sqlite3')
        self.assertEqual(moved.execute('SELECT key FROM ocr').fetchall(), [('page',)])
        moved.close()
        self.assertEqual(reporting.provenance_path(self.root).read_text(), '{"name": "1_2"}\n')
        self.assertTrue(reporting.registry_path(self.root).is_file())
        self.assertEqual(sorted(path.name for path in legacy.iterdir()), [])

    def test_existing_new_state_is_never_overwritten(self):
        (self.root / 'Reports').mkdir()
        (self.root / 'Reports' / 'review_assignments.jsonl').write_text('old\n')
        internal = self.root / 'Rapports' / '.interne'
        internal.mkdir(parents=True)
        (internal / 'review_assignments.jsonl').write_text('new\n')
        reporting.prepare(self.root)
        self.assertEqual(reporting.provenance_path(self.root).read_text(), 'new\n')
        self.assertTrue((self.root / 'Reports' / 'review_assignments.jsonl').exists())

    def test_execution_reports_never_overwrite_each_other(self):
        moment = datetime(2026, 10, 13, 9, 5, 7)
        paths = [reporting.execution_report_path(self.root, moment, 'ocr') for _ in range(2)]
        self.assertEqual([path.name for path in paths],
                         ['2026-10-13_09-05-07 - ocr.txt', '2026-10-13_09-05-07 - ocr (2).txt'])

    def test_command_line_rebuilds_every_report(self):
        with Store(self.root) as store:
            with store.replace_lot('7', 'run', [scanned('7', 1, VIS)]):
                pass
            reporting.refresh(self.root, store)
        (self.root / 'Rapports' / 'Lots' / '7' / 'pages.xlsx').unlink()
        self.assertEqual(reporting.main(['-d', str(self.root)]), 1)   # no search list known yet
        self.assertFalse((self.root / 'Rapports' / 'Lots' / '7' / 'pages.xlsx').exists())
        write_list(self.root / 'list.xlsx', [['1', VIS]])
        self.assertEqual(reporting.main(['-d', str(self.root), '--excel', 'list.xlsx', '--tout']), 0)
        self.assertTrue((self.root / 'Rapports' / 'Lots' / '7' / 'pages.xlsx').is_file())
        self.assertTrue((self.root / 'Rapports' / 'list_found.xlsx').is_file())

    def test_page_names_identify_lot_and_page(self):
        self.assertEqual(reporting.page_identity('520_14'), ('520', 14))
        self.assertIsNone(reporting.page_identity('scan'))
        self.assertEqual(reporting.relative(self.root, self.root / 'Output' / VIS / '1_1.pdf'), f'Output/{VIS}/1_1.pdf')
        self.assertEqual(reporting.relative(self.root, Path('/elsewhere/1_1.pdf')), '/elsewhere/1_1.pdf')


class ReviewStateTests(ReportingTestCase):
    def test_final_classification_moves_pages_to_review_or_classified(self):
        self.scan('3', 2)
        for page, state in ((1, CLASSIFIED), (2, TO_REVIEW)):
            vis = VIS if state == CLASSIFIED else ''
            self.store.record('review', 'run', PageEntry('3', page, state, (f'3_{page}', '', '', vis, '', '', ''),
                                                         'PREVIOUS_PAGE' if vis else 'NOT_SPARSE', seconds=1.5))
        reporting.refresh(self.root, self.store)
        rows = self.global_pages()
        self.assertEqual([(row['État'], row['Classée par'], row['Classification']) for row in rows],
                         [('Classée', 'Classification finale', 'PREVIOUS_PAGE'), ('À revoir', None, 'NOT_SPARSE')])
        [lot] = sheet_rows(self.root / 'Rapports' / 'suivi_lots.xlsx')
        self.assertEqual((lot['Durée classification (cumulée)'], lot['Durée code-barres'], lot['État']),
                         ('00:00:03', '00:00:12', 'Terminé'))


if __name__ == '__main__':
    unittest.main()
