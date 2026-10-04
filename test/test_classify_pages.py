from __future__ import annotations

import csv
import itertools
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace, ModuleType
from unittest.mock import patch

import main
from scripts import classify_pages as classify
from scripts import script as barcode
from test_split_pages import make_pdf, page_count, page_text


class ResolutionTests(unittest.TestCase):
    def test_publish_lot_ignores_appledouble_sidecars_on_removable_drives(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            work, output = root / 'work', root / 'output'
            (work / 'OCR').mkdir(parents=True)
            (work / '_reports').mkdir()
            (output / 'OCR').mkdir(parents=True)
            (output / 'lot_reports').mkdir()
            (output / '._OCR').write_bytes(b'existing metadata')
            (work / '._OCR').write_bytes(b'directory metadata')
            (work / 'OCR' / '1_1.pdf').write_bytes(b'real PDF')
            (work / 'OCR' / '._1_1.pdf').write_bytes(b'PDF metadata')
            (work / 'OCR' / '.DS_Store').write_bytes(b'Finder metadata')
            (work / '._reports').write_bytes(b'report metadata')
            classify.publish_lot(work, output, '1')
            self.assertEqual((output / 'OCR' / '1_1.pdf').read_bytes(), b'real PDF')
            self.assertFalse((output / 'OCR' / '._1_1.pdf').exists())
            self.assertFalse(work.exists())

    def test_report_enrichment_preserves_zeroes_and_leaves_unknowns_blank(self):
        record = classify.VehicleRecord('00123', 'VF1ABCDEFT5702376', 'T5702376',
                                        'SEQEMON0100000123', '4A2D0004')
        db = classify.DatabaseIndex({'00123': frozenset({'T5702376'})}, {'T5702376': [record]})
        expected = ('1_2', '00123', 'VF1ABCDEFT5702376', 'T5702376',
                    'SEQEMON0100000123', '100000123', '4A2D0004')
        for identifiers in ({'badges': ('00123',)}, {'vins': ('VF1ABCDEFT5702376',)}):
            self.assertEqual(classify.page_excel_row(barcode.PageResult('1', 2, 3, **identifiers),
                                                    'T5702376', db), expected)
        page = barcode.PageResult('1', 2, 3, badges=('99999',))
        self.assertEqual(classify.page_excel_row(page, None, db), ('1_2', '', '', '', '', '', ''))
        # A decoded VIN absent from the BDD stays usable; absent fields stay blank.
        page = barcode.PageResult('1', 2, 3, vins=('VF1ABCDEFT5702377',))
        self.assertEqual(classify.page_excel_row(page, 'T5702377', db),
                         ('1_2', '', 'VF1ABCDEFT5702377', 'T5702377', '', '', ''))

    def test_report_does_not_choose_between_conflicting_database_values(self):
        db = classify.DatabaseIndex(records_by_vis={'T5702376': [
            classify.VehicleRecord('123', '', 'T5702376', 'SEQ000000001', 'NOF1'),
            classify.VehicleRecord('456', '', 'T5702376', 'SEQ000000002', 'NOF2'),
        ]})
        page = barcode.PageResult('1', 1, 1, vins=('VF1ABCDEFT5702376',))
        self.assertEqual(classify.page_excel_row(page, 'T5702376', db),
                         ('1_1', '', 'VF1ABCDEFT5702376', 'T5702376', '', '', ''))
        page = barcode.PageResult('1', 1, 1, badges=('123',))
        self.assertEqual(classify.page_excel_row(page, 'T5702376', db),
                         ('1_1', '123', '', 'T5702376', 'SEQ000000001', '000000001', 'NOF1'))

    def test_bdd_extracts_vin_from_id2_and_seq_from_sem(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'BDD.xlsx'
            wb = Workbook()
            wb.active.append(['BDG', 'VIS', 'SEM', 'ID2', 'NOF'])
            wb.active.append(['00123', 'T5702376', '50FV750960305', 'VF1ABCDEFT5702376', '4A2D0004'])
            wb.active.append(['456', 'T5702377', '', 'not a VIN', ''])
            wb.save(path)
            wb.close()
            db = classify.load_database(path)
            page = barcode.PageResult('1', 1, 1, badges=('00123',))
            self.assertEqual(classify.page_excel_row(page, 'T5702376', db),
                             ('1_1', '00123', 'VF1ABCDEFT5702376', 'T5702376',
                              '50FV750960305', '750960305', '4A2D0004'))
            self.assertEqual(db.records_by_vis['T5702377'][0].vin, '')

    def test_badge_fallback_and_direct_vin(self):
        db = {'100000': frozenset({'T5708495'})}
        page = barcode.PageResult('479', 1, 2, badges=('100000',))
        self.assertEqual(classify.resolve_vis(page, db), ('T5708495', 'BDD_BADGE'))
        page = barcode.PageResult('479', 1, 2, vins=('VF1ABCDEFT5702376',))
        self.assertEqual(classify.resolve_vis(page, db), ('T5702376', 'BARCODE'))

    def test_conflicting_or_unresolved_identifiers_do_not_choose_a_vis(self):
        cases = [
            (barcode.PageResult('1', 1, 1), {}, 'NO_BARCODE'),
            (barcode.PageResult('1', 1, 1, badges=('12',)), {}, 'UNRESOLVED_BADGE'),
            (barcode.PageResult('1', 1, 1, error='broken'), {}, 'SCAN_ERROR'),
            (barcode.PageResult('1', 1, 1, badges=('12',)),
             {'12': frozenset({'T5702376', 'T5702377'})}, 'BDD_CONFLICT'),
            (barcode.PageResult('1', 1, 1, badges=('12',), vins=('VF1ABCDEFT5702376',)),
             {'12': frozenset({'T5702377'})}, 'IDENTIFIER_CONFLICT'),
            (barcode.PageResult('1', 1, 1, vins=('VF1ABCDEFT5702376', 'VF1ABCDEFT5702377')),
             {}, 'IDENTIFIER_CONFLICT'),
        ]
        for page, db, status in cases:
            with self.subTest(status=status):
                self.assertEqual(classify.resolve_vis(page, db), (None, status))

    def test_database_streams_headers_duplicates_and_padded_badges(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'database.xlsx'
            wb = Workbook()
            sheet = wb.active
            sheet.append(['Année', 'BDG', 'VIS'])
            sheet.append([2026, 100000, 'T5708495'])
            sheet.append([2025, 100000, 'T5708495'])
            sheet.append([2024, 12, 'T5702376'])
            sheet['B4'].number_format = '00000'
            sheet.append([2025, '0012', 'T5702377'])
            sheet.append([2026, '0012', 'T5702378'])
            wb.save(path)
            wb.close()
            db = classify.load_database(path).badge_vis
            self.assertEqual(db['100000'], frozenset({'T5708495'}))
            self.assertEqual(db['00012'], frozenset({'T5702376'}))
            self.assertEqual(db['0012'], frozenset({'T5702377', 'T5702378'}))

    def test_six_digit_badges_are_opt_in_and_direct_vis_is_ignored(self):
        module = ModuleType('codarascan')
        class Symbol:
            def __init__(self, text):
                self.text = text
        module.DecodedSymbolResult = Symbol
        scanner = SimpleNamespace(scan_image=lambda *a, **k: SimpleNamespace(
            metadata=SimpleNamespace(mode=barcode.PANORAMA_MODE, engine=barcode.PANORAMA_ENGINE),
            symbols=[Symbol('100000'), Symbol('T5708495')]))
        with (patch.dict('sys.modules', {'codarascan': module}),
              patch.object(barcode, '_WORKER_SCANNER', scanner),
              patch.object(barcode, '_worker_document', return_value=object()),
              patch.object(barcode, 'render_upper_half', return_value=SimpleNamespace(close=lambda: None))):
            original = barcode.process_page(barcode.PageTask('1', 'page.pdf', 0, 1))
            extended = barcode.process_page(barcode.PageTask('1', 'page.pdf', 0, 1, True))
        self.assertEqual(original.badges, ())
        self.assertEqual(original.vins, ())
        self.assertEqual(extended.badges, ('100000',))
        self.assertEqual(extended.vins, ())
        self.assertEqual(classify.resolve_vis(extended, {}), (None, 'UNRESOLVED_BADGE'))

    def test_nof_barcode_does_not_conflict_with_vin_on_page_1_6(self):
        module = ModuleType('codarascan')
        class Symbol:
            def __init__(self, text):
                self.text = text
        module.DecodedSymbolResult = Symbol
        scanner = SimpleNamespace(scan_image=lambda *a, **k: SimpleNamespace(
            metadata=SimpleNamespace(mode=barcode.PANORAMA_MODE, engine=barcode.PANORAMA_ENGINE),
            symbols=[Symbol('BRYEKNFJXT5704852'), Symbol('6G3D1365'), Symbol('T5702376')]))
        with (patch.dict('sys.modules', {'codarascan': module}),
              patch.object(barcode, '_WORKER_SCANNER', scanner),
              patch.object(barcode, '_worker_document', return_value=object()),
              patch.object(barcode, 'render_upper_half', return_value=SimpleNamespace(close=lambda: None))):
            result = barcode.process_page(barcode.PageTask('1', 'page.pdf', 0, 1, True))
        self.assertEqual(result.vins, ('BRYEKNFJXT5704852',))
        self.assertEqual(classify.resolve_vis(result, {}), ('T5704852', 'BARCODE'))


class WorkflowTests(unittest.TestCase):
    @staticmethod
    def prepare_failure_batch(root):
        from openpyxl import Workbook
        for lot in ('1', '2', '3'):
            (root / lot).mkdir()
            make_pdf(root / lot / 'lot.pdf', [f'{lot} page 1', f'{lot} page 2'])
        for name in ('BDD.xlsx', 'search.xlsx'):
            wb = Workbook()
            wb.active.append(['BDG', 'VIS'])
            wb.active.append(['12345', 'T5708495'])
            wb.save(root / name)
            wb.close()
        class Pool:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def imap_unordered(self, function, tasks, chunksize):
                return (function(task) for task in tasks)
        return classify.parse_args(['-d', str(root)]), SimpleNamespace(Pool=lambda *a, **k: Pool())

    @staticmethod
    def decoded_page(task):
        return barcode.PageResult(task.lot, 1, task.pages_in_lot, vins=('VF1ABCDEFT5708495',))

    def test_later_lot_exception_preserves_previous_lots_and_continues(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args, context = self.prepare_failure_batch(root)
            def scan(task):
                if Path(task.pdf_path).name == '2_2.pdf':
                    raise RuntimeError('deliberate worker failure')
                return self.decoded_page(task)
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(barcode, 'process_page', side_effect=scan),
                  patch.object(barcode, '_close_worker_document'),
                  patch.object(classify.mp, 'get_context', return_value=context)):
                self.assertEqual(classify.run(args), 1)
            output = root / 'output'
            self.assertTrue((output / 'T5708495/1_1.pdf').is_file())
            self.assertTrue((output / 'T5708495/3_2.pdf').is_file())
            self.assertFalse((output / 'T5708495/2_1.pdf').exists())
            self.assertTrue(list((output / '_incomplete').glob('2_*/T5708495/2_1.pdf')))
            pages = load_workbook(output / 'pages.xlsx', read_only=True)
            self.assertEqual([row[0] for row in list(pages.active.values)[1:]], ['1_1', '1_2', '3_1', '3_2'])
            pages.close()
            report = (output / 'report.txt').read_text()
            self.assertIn('Batch status: COMPLETED_WITH_ISSUES', report)
            self.assertIn('Lot 2: FAILED', report)
            self.assertIn('Lots published: 2', report)

    def test_interrupt_keeps_completed_lot_pdfs_csv_excel_and_timing(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args, context = self.prepare_failure_batch(root)
            def scan(task):
                if Path(task.pdf_path).name == '2_2.pdf':
                    raise KeyboardInterrupt()
                return self.decoded_page(task)
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(barcode, 'process_page', side_effect=scan),
                  patch.object(barcode, '_close_worker_document'),
                  patch.object(classify.mp, 'get_context', return_value=context)):
                with self.assertRaises(KeyboardInterrupt):
                    classify.run(args)
            output = root / 'output'
            self.assertEqual(page_text(output / 'T5708495/1_1.pdf'), '1 page 1')
            self.assertEqual(page_text(output / 'T5708495/1_2.pdf'), '1 page 2')
            with (output / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 2)
            pages = load_workbook(output / 'lot_reports/1/pages.xlsx', read_only=True)
            self.assertEqual([row[0] for row in list(pages.active.values)[1:]], ['1_1', '1_2'])
            pages.close()
            self.assertIn('Batch status: INTERRUPTED', (output / 'report.txt').read_text())
            self.assertIn('Lot 1: SAVED', (output / 'report.txt').read_text())
            self.assertIn('Lot 2: INTERRUPTED', (output / 'report.txt').read_text())
            self.assertIn('Lot 1 | SAVED | duration=', (output / 'processing.log').read_text())
            self.assertTrue(list((output / '_incomplete').glob('2_*')))

    def test_final_excel_export_failure_keeps_all_published_lots(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args, context = self.prepare_failure_batch(root)
            original_export = classify.write_page_workbook
            def export(path, rows):
                if path.name == '.pages.xlsx.tmp':
                    raise OSError('deliberate final export failure')
                return original_export(path, rows)
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(barcode, 'process_page', side_effect=self.decoded_page),
                  patch.object(barcode, '_close_worker_document'),
                  patch.object(classify.mp, 'get_context', return_value=context),
                  patch.object(classify, 'write_page_workbook', side_effect=export)):
                with self.assertRaisesRegex(OSError, 'final export'):
                    classify.run(args)
            output = root / 'output'
            for lot in ('1', '2', '3'):
                self.assertTrue((output / f'T5708495/{lot}_2.pdf').is_file())
                self.assertTrue((output / f'lot_reports/{lot}/pages.xlsx').is_file())
            with (output / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 6)
            self.assertIn('Batch status: FAILED', (output / 'report.txt').read_text())
            self.assertIn('Lots published: 3', (output / 'report.txt').read_text())

    def test_publish_failure_keeps_saved_lots_and_current_lot_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            args, context = self.prepare_failure_batch(root)
            original_rename = Path.rename
            def rename(path, target):
                if Path(target).resolve() == (root / 'output/T5708495/2_2.pdf').resolve():
                    raise OSError('deliberate publication failure')
                return original_rename(path, target)
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(barcode, 'process_page', side_effect=self.decoded_page),
                  patch.object(barcode, '_close_worker_document'),
                  patch.object(classify.mp, 'get_context', return_value=context),
                  patch.object(Path, 'rename', rename)):
                self.assertEqual(classify.run(args), 1)
            output = root / 'output'
            self.assertTrue((output / 'T5708495/1_2.pdf').is_file())
            self.assertTrue((output / 'T5708495/2_1.pdf').is_file())
            self.assertTrue((output / 'T5708495/3_2.pdf').is_file())
            self.assertTrue(list((output / '_incomplete').glob('2_*/T5708495/2_2.pdf')))
            self.assertTrue(list((output / '_incomplete').glob('2_*/_reports/results.csv')))
            self.assertTrue(list((output / '_incomplete').glob('2_*/_reports/pages.xlsx')))
            self.assertIn('Lot 2: FAILED', (output / 'report.txt').read_text())

    def test_database_discovery_ignores_search_list_and_temporary_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name in ('BDD_2024_2025_2026.xlsx', 'Dirty List VIN Delivered.xlsx',
                         '._BDD_2024_2025_2026.xlsx', '~$BDD_2024_2025_2026.xlsx',
                         'BDD_2024_2025_2026_found.xlsx'):
                (root / name).touch()
            self.assertEqual(classify.discover_database(root),
                             (root / 'BDD_2024_2025_2026.xlsx').resolve())
            self.assertIsNone(classify.parse_args(['-d', str(root)]).database)
            (root / 'bdd24-26.xlsx').touch()
            with self.assertRaisesRegex(barcode.ConfigurationError, 'Multiple BDD'):
                classify.discover_database(root)
            self.assertEqual(classify.discover_database(root, Path('bdd24-26.xlsx')),
                             (root / 'bdd24-26.xlsx').resolve())

    def test_missing_database_has_clear_error(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(barcode.ConfigurationError, 'No BDD workbook'):
                classify.discover_database(Path(directory))

    def test_menu_command_allows_database_auto_discovery(self):
        command = main.build_command('classify_pages.py', Path('/batch'), workers=2)
        self.assertNotIn('--database', command)

    def test_end_to_end_classification_search_list_and_backup(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for lot, labels in [('479', ['badge only', 'unresolved', 'scan error']),
                                ('500', ['same VIS', 'conflict'])]:
                (root / lot).mkdir()
                make_pdf(root / lot / 'lot.pdf', labels)
            (root / 'notes').mkdir()
            (root / '501').mkdir()  # Must be reported, not silently ignored.
            (root / 'output').mkdir()
            (root / 'output' / 'old.txt').write_text('preserve')
            db_path = root / 'BDD.xlsx'
            target_path = root / 'search.xlsx'
            for path, rows in [(db_path, [['BDG', 'VIS', 'SEM', 'ID2', 'NOF'],
                                         ['100000', 'T5708495', 'SEQEMON0117700049', 'VF1ABCDEFT5708495', '4A2D0004']]),
                               (target_path, [['BDG', 'VIS'], ['12345', 'T5708495'],
                                              ['54321', 'T5702376']])]:
                wb = Workbook()
                for row in rows:
                    wb.active.append(row)
                wb.save(path)
                wb.close()
            originals = {p: p.read_bytes() for p in root.glob('*/*.pdf')}
            scanned = []
            def scan(task):
                self.assertEqual(task.page_index, 0)
                self.assertTrue(task.six_digit_badges)
                self.assertEqual(page_count(Path(task.pdf_path)), 1)
                name = Path(task.pdf_path).name
                output = root / 'output'
                if task.lot == '479':
                    self.assertFalse(list(output.glob('T5708495/*.pdf')))
                    self.assertFalse(list((output / 'OCR').glob('*.pdf')))
                elif task.lot == '500':
                    self.assertTrue((output / 'T5708495/479_1.pdf').is_file())
                    self.assertTrue((output / 'OCR/479_2.pdf').is_file())
                    self.assertTrue((output / 'lot_reports/479/pages.xlsx').is_file())
                    with (output / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                        saved = list(csv.DictReader(handle))
                    self.assertEqual(len(saved), 3)
                    self.assertTrue(all(row['Lot'] == '479' for row in saved))
                    self.assertIn('Lots published: 1', (output / 'report.txt').read_text())
                    self.assertFalse((output / 'T5708495/500_1.pdf').exists())
                scanned.append(name)
                values = {
                    '479_1.pdf': {'badges': ('100000',)},
                    '479_2.pdf': {},
                    '479_3.pdf': {'error': 'deliberate scan failure'},
                    '500_1.pdf': {'vins': ('VF1ABCDEFT5708495',)},
                    '500_2.pdf': {'badges': ('100000',), 'vins': ('VF1ABCDEFT5702376',)},
                }[name]
                return barcode.PageResult(task.lot, 1, task.pages_in_lot, **values)
            class Pool:
                def __enter__(self): return self
                def __exit__(self, *args): pass
                def imap_unordered(self, function, tasks, chunksize):
                    return (function(task) for task in reversed(tasks))
            args = classify.parse_args(['-d', str(root)])
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(barcode, 'process_page', side_effect=scan),
                  patch.object(barcode, '_close_worker_document'),
                  patch.object(classify.time, 'perf_counter', side_effect=itertools.count(100.0, 10.0)) as clock,
                  patch.object(classify, 'load_database', wraps=classify.load_database) as load_bdd,
                  patch.object(classify.mp, 'get_context', return_value=SimpleNamespace(Pool=lambda *a, **k: Pool()))):
                original_loader = load_bdd._mock_wraps
                def timed_loader(path):
                    self.assertEqual(clock.call_count, 1)  # Timing starts before BDD loading.
                    return original_loader(path)
                load_bdd.side_effect = timed_loader
                code = classify.run(args)
            self.assertEqual(code, 1)
            output = root / 'output'
            self.assertEqual(page_text(output / 'T5708495/479_1.pdf'), 'badge only')
            self.assertEqual(page_text(output / 'T5708495/500_1.pdf'), 'same VIS')
            self.assertEqual(len(list((output / 'OCR').glob('*.pdf'))), 3)
            self.assertEqual(len(scanned), 5)
            with (output / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([(r['Lot'], r['Page']) for r in rows],
                             [('479', '1'), ('479', '2'), ('479', '3'), ('500', '1'), ('500', '2')])
            self.assertEqual(rows[0]['VIS'], 'T5708495')
            self.assertEqual(rows[0]['Decoded VIS'], '')
            self.assertEqual(rows[0]['Found'], 'FOUND')
            self.assertEqual(rows[2]['Status'], 'SCAN_ERROR')
            self.assertEqual(rows[-1]['Status'], 'IDENTIFIER_CONFLICT')
            result = load_workbook(output / 'search_found.xlsx')
            self.assertEqual(result.active['A2'].fill.fgColor.rgb, '00C6EFCE')
            result.close()
            pages = load_workbook(output / 'pages.xlsx')
            sheet = pages.active
            page_rows = list(sheet.values)
            self.assertEqual(page_rows[0], ('lot_page', 'BDG', 'VIN', 'VIS', 'SEQ', 'SEQ_9', 'NOF'))
            self.assertEqual(page_rows[1], ('479_1', '100000', 'VF1ABCDEFT5708495', 'T5708495',
                                           'SEQEMON0117700049', '117700049', '4A2D0004'))
            self.assertEqual(page_rows[4][1:], page_rows[1][1:])
            for index in (2, 3, 5):
                self.assertEqual(page_rows[index][1:], (None,) * 6)
            self.assertEqual(len(page_rows), 6)
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertEqual(sheet.auto_filter.ref, 'A1:G6')
            self.assertEqual(sheet['F2'].number_format, '@')
            pages.close()
            self.assertIn('Lot 501: no PDF found', (output / 'report.txt').read_text())
            report = (output / 'report.txt').read_text()
            self.assertIn('Start time:', report)
            self.assertIn('Finish time:', report)
            self.assertIn('Total processing time: 00:01:50', report)
            self.assertIn('Total processing seconds: 110.000', report)
            log = (output / 'processing.log').read_text()
            self.assertIn('Batch started:', log)
            self.assertIn('Batch finished:', log)
            self.assertIn('00:01:50 (110.000 seconds)', log)
            self.assertIn('Lot 479 | START', log)
            self.assertIn('Lot 479 | SAVED_WITH_ISSUES | duration=00:00:10 (10.000 seconds)', log)
            with (output / 'lots.csv').open(encoding='utf-8-sig', newline='') as handle:
                timing = list(csv.DictReader(handle))
            self.assertEqual([row['Lot'] for row in timing], ['479', '500', '501'])
            self.assertEqual([row['Seconds'] for row in timing], ['10.000'] * 3)
            self.assertEqual(timing[-1]['Status'], 'FAILED')
            self.assertEqual(len(list(root.glob('output_backup_*/old.txt'))), 1)
            self.assertFalse(list(root.glob('.classify_pages_*')))
            for path, content in originals.items():
                self.assertEqual(path.read_bytes(), content)

    def test_menu_builds_new_command_with_separate_database(self):
        command = main.build_command('classify_pages.py', Path('/batch'), workers=3,
                                     excel=Path('/batch/list.xlsx'), database=Path('/BDD.xlsx'))
        self.assertEqual(command[-6:], ['-n', '3', '--excel', '/batch/list.xlsx', '--database', '/BDD.xlsx'])

    def test_fatal_split_failure_preserves_previous_output_in_backup(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / '1').mkdir()
            (root / '1/broken.pdf').write_bytes(b'not a PDF')
            (root / 'output').mkdir()
            (root / 'output/old.txt').write_text('keep')
            for name in ('BDD.xlsx', 'search.xlsx'):
                wb = Workbook()
                wb.active.append(['BDG', 'VIS'])
                wb.active.append(['12345', 'T5702376'])
                wb.save(root / name)
                wb.close()
            class Pool:
                def __enter__(self): return self
                def __exit__(self, *args): pass
            args = classify.parse_args(['-d', str(root), '--database', 'BDD.xlsx'])
            with (patch.object(barcode, 'validate_runtime'),
                  patch.object(classify.mp, 'get_context', return_value=SimpleNamespace(Pool=lambda *a, **k: Pool()))):
                with self.assertRaises(barcode.ConfigurationError):
                    classify.run(args)
            backups = list(root.glob('output_backup_*/old.txt'))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_text(), 'keep')
            self.assertIn('Batch status: FAILED', (root / 'output/report.txt').read_text())
            self.assertTrue(list((root / 'output/_incomplete').glob('1_*')))
            self.assertFalse(list(root.glob('.classify_pages_*')))


if __name__ == '__main__':
    unittest.main()
